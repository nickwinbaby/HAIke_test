import ctypes
import json
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill


# ====== 可直接在这里修改 ======
MCUKEY = "3F17FC9852ACC2856C57805589AE5A363B3522F48378398F80686D00EEA7ACB6"
DLL_PATH = r".\dll\jsonrichlinksdk.dll"
EXCEL_PATH = r"gmy.xlsx"
# ==============================


class JsonLogLevel:
    JSON_LOG_LEVEL_DEBUG = 0
    JSON_LOG_LEVEL_INFO = 1
    JSON_LOG_LEVEL_WARN = 2
    JSON_LOG_LEVEL_ERROR = 3


PROTOCOL_VERSIONS = {
    0: "1.1.0",
    1: "1.1.1",
    2: "2.5.1",
    3: "2.5.2",
    4: "G - 不支持",
    5: "4.0.0",
    6: "4.0.1",
    7: "4.0.2",
}


OUT_BUFFER_SIZE = 3072
MODE_JSON_TO_HEX = "JSON→HEX"
MODE_HEX_TO_JSON = "HEX→JSON"


def key_to_bytes(mcukey):
    return bytes.fromhex(mcukey.strip())


def normalize_hex(value):
    if value is None:
        return ""
    return "".join(str(value).split()).lower()


def normalize_json_text(value):
    if value is None:
        return ""
    parsed_json = json.loads(str(value))
    return json.dumps(parsed_json, indent=4)


class JsonSdk:
    def __init__(self, dll_path, protocol_version):
        self.protocol_version = protocol_version
        self.json_sdk = ctypes.CDLL(str(dll_path))
        self.json_sdk.json_protocol_get_library_version.restype = ctypes.c_char_p
        self.json_sdk.json_protocol_get_error_description.argtypes = [ctypes.c_int]
        self.json_sdk.json_protocol_get_error_description.restype = ctypes.c_char_p
        self.json_sdk.json_protocol_init_logger.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_char_p]
        self.json_sdk.json_protocol_frame_package.argtypes = [
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.POINTER(ctypes.c_int),
            ctypes.c_char_p,
            ctypes.c_int,
        ]
        self.json_sdk.json_protocol_frame_package.restype = ctypes.c_int
        self.json_sdk.json_protocol_frame_parse.argtypes = [
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_int,
        ]
        self.json_sdk.json_protocol_frame_parse.restype = ctypes.c_int

    def get_library_version(self):
        version = self.json_sdk.json_protocol_get_library_version()
        return version.decode("utf-8", errors="replace")

    def get_error_description(self, err_code):
        description = self.json_sdk.json_protocol_get_error_description(err_code)
        if not description:
            return ""
        return description.decode("utf-8", errors="replace")

    def initialise(self, log_dir):
        os.makedirs(log_dir, exist_ok=True)
        self.json_sdk.json_protocol_init_logger(
            JsonLogLevel.JSON_LOG_LEVEL_DEBUG,
            str(log_dir).encode("utf-8"),
            b"logfile.log",
        )
        print("Library Version:", self.get_library_version())
        print("Error Description (Code 0):", self.get_error_description(0))

    def frame_parse(self, in_buffer, in_length, out_buffer, out_length, mcukey):
        out_length_value = ctypes.c_int(out_length)
        result = self.json_sdk.json_protocol_frame_parse(
            in_buffer,
            in_length,
            out_buffer,
            out_length_value,
            key_to_bytes(mcukey),
            self.protocol_version,
        )
        parsed_data = out_buffer.raw[: out_length_value.value].decode("utf-8", errors="replace").rstrip("\x00")
        return result, parsed_data

    def frame_package(self, json_text, out_buffer, out_length, mcukey):
        out_length_value = ctypes.c_int(out_length)
        json_text = str(json_text)
        result = self.json_sdk.json_protocol_frame_package(
            json_text.encode("utf-8"),
            len(json_text),
            out_buffer,
            ctypes.byref(out_length_value),
            key_to_bytes(mcukey),
            self.protocol_version,
        )
        return result, out_buffer.raw[: out_length_value.value]


def parse_hex_to_json(json_sdk, hex_text, mcukey):
    normalized_hex = normalize_hex(hex_text)
    in_buffer = ctypes.create_string_buffer(bytes.fromhex(normalized_hex))
    in_length = len(normalized_hex) // 2
    out_buffer = ctypes.create_string_buffer(OUT_BUFFER_SIZE)
    ctypes.memset(out_buffer, 0, OUT_BUFFER_SIZE)

    parse_result, parsed_data = json_sdk.frame_parse(
        in_buffer,
        in_length,
        out_buffer,
        OUT_BUFFER_SIZE,
        mcukey,
    )
    if parse_result != 0:
        raise RuntimeError(f"{parse_result}: {json_sdk.get_error_description(parse_result)}")

    parsed_json = json.loads(parsed_data)
    return json.dumps(parsed_json, indent=4)


def package_json_to_hex(json_sdk, json_text, mcukey):
    if json_text is None:
        raise ValueError("JSON 内容为空")
    out_buffer = ctypes.create_string_buffer(OUT_BUFFER_SIZE)
    ctypes.memset(out_buffer, 0, OUT_BUFFER_SIZE)
    package_result, packaged_data = json_sdk.frame_package(json_text, out_buffer, OUT_BUFFER_SIZE, mcukey)
    if package_result != 0:
        raise RuntimeError(f"{package_result}: {json_sdk.get_error_description(package_result)}")
    return packaged_data.hex()


def build_output_path(excel_path, conversion_mode):
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    mode_suffix = "JSON-HEX" if conversion_mode == MODE_JSON_TO_HEX else "HEX-JSON"
    return excel_path.with_name(f"{excel_path.stem}-{mode_suffix}-{timestamp}{excel_path.suffix}")


def make_stats():
    return {
        "first_parse_success": 0,
        "first_parse_fail": 0,
        "second_parse_success": 0,
        "second_parse_fail": 0,
    }


def fill_result(cell, result, red_fill, green_fill):
    cell.value = result
    cell.fill = green_fill if result == "PASS" else red_fill


def prepare_workbook(dll_path, excel_path, output_path, mcukey, protocol_version):
    json_sdk = JsonSdk(dll_path, protocol_version)
    json_sdk.initialise(output_path.parent / "logs")

    workbook = openpyxl.load_workbook(excel_path)
    sheet = workbook.active

    protocol_description = PROTOCOL_VERSIONS.get(protocol_version, "未知协议")
    now_ms = int(datetime.now(timezone.utc).timestamp() * 1000)
    bjtime = (datetime.utcfromtimestamp(now_ms / 1000) + timedelta(hours=8)).strftime("%Y-%m-%d %H:%M:%S")

    sheet.cell(row=1, column=10, value="MCUKEY：" + mcukey)
    sheet.cell(row=2, column=10, value="执行时间：" + bjtime)
    sheet.cell(row=3, column=10, value="协议版本号：" + protocol_description)

    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    green_fill = PatternFill(start_color="FF00B050", end_color="FF00B050", fill_type="solid")
    return json_sdk, workbook, sheet, red_fill, green_fill


def process_hex_to_json(dll_path, excel_path, output_path, mcukey, protocol_version):
    json_sdk, workbook, sheet, red_fill, green_fill = prepare_workbook(
        dll_path, excel_path, output_path, mcukey, protocol_version
    )
    stats = make_stats()

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=1, values_only=True), start=1):
        in_buffer_hex = row[0]
        if in_buffer_hex is None:
            continue

        try:
            formatted_json = parse_hex_to_json(json_sdk, in_buffer_hex, mcukey)
            sheet.cell(row=row_idx, column=2, value=formatted_json)
            formatted_hex = package_json_to_hex(json_sdk, formatted_json, mcukey)
            sheet.cell(row=row_idx, column=3, value=formatted_hex)
            stats["first_parse_success"] += 1
            print(f"第 {row_idx} 行 A->B/C 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=2, value=f"data fail: {exc}")
            stats["first_parse_fail"] += 1
            print(f"第 {row_idx} 行 A->B/C 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=3, values_only=True), start=1):
        a_col_value = normalize_hex(row[0])
        c_col_value = normalize_hex(row[2])
        result = "PASS" if c_col_value and a_col_value == c_col_value else "FAIL"
        fill_result(sheet.cell(row=row_idx, column=4), result, red_fill, green_fill)

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, min_col=3, max_col=3, values_only=True), start=1):
        in_buffer_hex = row[0]
        if in_buffer_hex is None:
            continue

        try:
            formatted_json = parse_hex_to_json(json_sdk, in_buffer_hex, mcukey)
            sheet.cell(row=row_idx, column=6, value=formatted_json)

            parsed_json = json.loads(formatted_json)
            msg_id_desc = parsed_json.get("body", {}).get("MsgIDDesc", "")
            sheet.cell(row=row_idx, column=8, value=msg_id_desc)

            stats["second_parse_success"] += 1
            print(f"第 {row_idx} 行 C->F 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=6, value=f"data fail: {exc}")
            stats["second_parse_fail"] += 1
            print(f"第 {row_idx} 行 C->F 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=6, values_only=True), start=1):
        b_col_value = row[1]
        f_col_value = row[5]
        result = "PASS" if b_col_value and b_col_value == f_col_value else "FAIL"
        fill_result(sheet.cell(row=row_idx, column=7), result, red_fill, green_fill)

    workbook.save(output_path)
    return stats


def process_json_to_hex(dll_path, excel_path, output_path, mcukey, protocol_version):
    json_sdk, workbook, sheet, red_fill, green_fill = prepare_workbook(
        dll_path, excel_path, output_path, mcukey, protocol_version
    )
    stats = make_stats()

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=1, values_only=True), start=1):
        in_json_data = row[0]
        if in_json_data is None:
            continue

        try:
            formatted_hex = package_json_to_hex(json_sdk, in_json_data, mcukey)
            sheet.cell(row=row_idx, column=2, value=formatted_hex)

            formatted_json = parse_hex_to_json(json_sdk, formatted_hex, mcukey)
            sheet.cell(row=row_idx, column=3, value=formatted_json)

            parsed_json = json.loads(formatted_json)
            msg_id_desc = parsed_json.get("body", {}).get("MsgIDDesc", "")
            sheet.cell(row=row_idx, column=8, value=msg_id_desc)

            stats["first_parse_success"] += 1
            print(f"第 {row_idx} 行 A->B/C 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=2, value=f"data fail: {exc}")
            stats["first_parse_fail"] += 1
            print(f"第 {row_idx} 行 A->B/C 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=3, values_only=True), start=1):
        a_col_value = row[0]
        c_col_value = row[2]
        try:
            result = "PASS" if normalize_json_text(a_col_value) == normalize_json_text(c_col_value) else "FAIL"
        except Exception:
            result = "FAIL"
        fill_result(sheet.cell(row=row_idx, column=4), result, red_fill, green_fill)

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, min_col=3, max_col=3, values_only=True), start=1):
        in_json_data = row[0]
        if in_json_data is None:
            continue

        try:
            formatted_hex = package_json_to_hex(json_sdk, in_json_data, mcukey)
            sheet.cell(row=row_idx, column=6, value=formatted_hex)
            stats["second_parse_success"] += 1
            print(f"第 {row_idx} 行 C->F 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=6, value=f"data fail: {exc}")
            stats["second_parse_fail"] += 1
            print(f"第 {row_idx} 行 C->F 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=6, values_only=True), start=1):
        b_col_value = normalize_hex(row[1])
        f_col_value = normalize_hex(row[5])
        result = "PASS" if b_col_value and b_col_value == f_col_value else "FAIL"
        fill_result(sheet.cell(row=row_idx, column=7), result, red_fill, green_fill)

    workbook.save(output_path)
    return stats


def choose_mode():
    print("请选择转换方向：")
    print("1. JSON→HEX")
    print("2. HEX→JSON")
    while True:
        choice = input("请输入 1 或 2：").strip()
        if choice == "1":
            return MODE_JSON_TO_HEX
        if choice == "2":
            return MODE_HEX_TO_JSON
        print("输入无效，请重新输入。")


def choose_protocol_version():
    print("\n请选择协议版本：")
    for protocol_id, description in PROTOCOL_VERSIONS.items():
        print(f"{protocol_id}. {description}")
    while True:
        choice = input("请输入协议版本序号：").strip()
        if choice.isdigit() and int(choice) in PROTOCOL_VERSIONS:
            return int(choice)
        print("协议版本序号无效，请重新输入。")


def validate_config(dll_path, excel_path, mcukey):
    if not dll_path.is_file():
        raise FileNotFoundError(f"DLL 文件不存在：{dll_path}")
    if not excel_path.is_file():
        raise FileNotFoundError(f"Excel 文件不存在：{excel_path}")
    if excel_path.suffix.lower() != ".xlsx":
        raise ValueError("当前只支持 .xlsx 文件。")
    key_to_bytes(mcukey)


def main():
    dll_path = Path(DLL_PATH)
    excel_path = Path(EXCEL_PATH)
    mcukey = MCUKEY

    conversion_mode = choose_mode()
    protocol_version = choose_protocol_version()
    validate_config(dll_path, excel_path, mcukey)

    output_path = build_output_path(excel_path, conversion_mode)

    print("\n开始处理")
    print(f"DLL: {dll_path}")
    print(f"Excel: {excel_path}")
    print(f"输出: {output_path}")
    print(f"转换方向: {conversion_mode}")
    print(f"协议: {protocol_version} ({PROTOCOL_VERSIONS[protocol_version]})")

    if conversion_mode == MODE_JSON_TO_HEX:
        stats = process_json_to_hex(dll_path, excel_path, output_path, mcukey, protocol_version)
    else:
        stats = process_hex_to_json(dll_path, excel_path, output_path, mcukey, protocol_version)

    print("\n处理完成")
    print(f"输出文件：{output_path}")
    print(f"A->B/C 成功：{stats['first_parse_success']}，失败：{stats['first_parse_fail']}")
    print(f"C->F 成功：{stats['second_parse_success']}，失败：{stats['second_parse_fail']}")


if __name__ == "__main__":
    main()
