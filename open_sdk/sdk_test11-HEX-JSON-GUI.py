import ctypes
import json
import os
import threading
from datetime import datetime, timezone, timedelta
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl
from openpyxl.styles import PatternFill


class JsonLogLevel:
    JSON_LOG_LEVEL_DEBUG = 0
    JSON_LOG_LEVEL_INFO = 1
    JSON_LOG_LEVEL_WARN = 2
    JSON_LOG_LEVEL_ERROR = 3


PROTOCOL_VERSIONS = {
    0: "1.1.0",
    1: "1.1.1",
    2: "2.5.1",
    3: "2.5.2",
    4: "G - 不支持",
    5: "4.0.0",
    6: "4.0.1",
    7: "4.0.2",
}


DEFAULT_MCUKEY = "3F17FC9852ACC2856C57805589AE5A363B3522F48378398F80686D00EEA7ACB6"
OUT_BUFFER_SIZE = 3072
MODE_HEX_TO_JSON = "HEX→JSON"
MODE_JSON_TO_HEX = "JSON→HEX"


def key_to_bytes(mcukey):
    return bytes.fromhex(mcukey.strip())


def normalize_hex(value):
    if value is None:
        return ""
    return "".join(str(value).split()).lower()


class JsonSdk:
    def __init__(self, dll_path, protocol_version):
        self.protocol_version = protocol_version
        self.json_sdk = ctypes.CDLL(str(dll_path))
        self.json_sdk.json_protocol_get_library_version.restype = ctypes.c_char_p
        self.json_sdk.json_protocol_get_error_description.argtypes = [ctypes.c_int]
        self.json_sdk.json_protocol_get_error_description.restype = ctypes.c_char_p
        self.json_sdk.json_protocol_init_logger.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_char_p]
        self.json_sdk.json_protocol_frame_package.argtypes = [
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.POINTER(ctypes.c_int),
            ctypes.c_char_p,
            ctypes.c_int,
        ]
        self.json_sdk.json_protocol_frame_package.restype = ctypes.c_int
        self.json_sdk.json_protocol_frame_parse.argtypes = [
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_int,
        ]
        self.json_sdk.json_protocol_frame_parse.restype = ctypes.c_int

    def get_library_version(self):
        version = self.json_sdk.json_protocol_get_library_version()
        return version.decode("utf-8", errors="replace")

    def get_error_description(self, err_code):
        description = self.json_sdk.json_protocol_get_error_description(err_code)
        if not description:
            return ""
        return description.decode("utf-8", errors="replace")

    def init_logger(self, level, output_path, file_name):
        self.json_sdk.json_protocol_init_logger(
            level,
            str(output_path).encode("utf-8"),
            file_name.encode("utf-8"),
        )

    def initialise(self, log_dir):
        os.makedirs(log_dir, exist_ok=True)
        self.init_logger(JsonLogLevel.JSON_LOG_LEVEL_DEBUG, log_dir, "logfile.log")
        return self.get_library_version(), self.get_error_description(0)

    def frame_parse(self, in_buffer, in_length, out_buffer, out_length, mcukey):
        out_length_value = ctypes.c_int(out_length)
        result = self.json_sdk.json_protocol_frame_parse(
            in_buffer,
            in_length,
            out_buffer,
            out_length_value,
            key_to_bytes(mcukey),
            self.protocol_version,
        )
        parsed_data = out_buffer.raw[: out_length_value.value].decode("utf-8", errors="replace").rstrip("\x00")
        return result, parsed_data

    def frame_package(self, json_text, out_buffer, out_length, mcukey):
        out_length_value = ctypes.c_int(out_length)
        result = self.json_sdk.json_protocol_frame_package(
            json_text.encode("utf-8"),
            len(json_text),
            out_buffer,
            ctypes.byref(out_length_value),
            key_to_bytes(mcukey),
            self.protocol_version,
        )
        return result, out_buffer.raw[: out_length_value.value]


class HexJsonGui(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("HEX-JSON 解析工具 -Mangetic-2026-8-13")
        self.geometry("820x600")
        self.minsize(760, 560)

        self.mode_var = tk.StringVar(value=MODE_HEX_TO_JSON)
        self.protocol_var = tk.StringVar(value=self.protocol_display(1))
        self.mcukey_var = tk.StringVar(value=DEFAULT_MCUKEY)
        self.dll_path_var = tk.StringVar()
        self.excel_path_var = tk.StringVar()
        self.status_var = tk.StringVar(value="请选择 DLL 和 Excel 文件")
        self.output_path = None
        self.worker = None

        self.create_widgets()

    @staticmethod
    def protocol_display(protocol_id):
        return f"{protocol_id}: {PROTOCOL_VERSIONS[protocol_id]}"

    def create_widgets(self):
        container = ttk.Frame(self, padding=16)
        container.pack(fill=tk.BOTH, expand=True)
        container.columnconfigure(1, weight=1)

        ttk.Label(container, text="转换方向").grid(row=0, column=0, sticky=tk.W, pady=6)
        mode_box = ttk.Combobox(
            container,
            textvariable=self.mode_var,
            values=[MODE_HEX_TO_JSON, MODE_JSON_TO_HEX],
            state="readonly",
            width=28,
        )
        mode_box.grid(row=0, column=1, sticky=tk.W, pady=6)

        ttk.Label(container, text="协议版本").grid(row=1, column=0, sticky=tk.W, pady=6)
        protocol_box = ttk.Combobox(
            container,
            textvariable=self.protocol_var,
            values=[self.protocol_display(i) for i in PROTOCOL_VERSIONS],
            state="readonly",
            width=28,
        )
        protocol_box.grid(row=1, column=1, sticky=tk.W, pady=6)

        ttk.Label(container, text="MCUKEY").grid(row=2, column=0, sticky=tk.W, pady=6)
        ttk.Entry(container, textvariable=self.mcukey_var).grid(row=2, column=1, sticky=tk.EW, pady=6)

        ttk.Label(container, text="DLL 文件").grid(row=3, column=0, sticky=tk.W, pady=6)
        ttk.Entry(container, textvariable=self.dll_path_var).grid(row=3, column=1, sticky=tk.EW, pady=6)
        ttk.Button(container, text="选择", command=self.choose_dll).grid(row=3, column=2, padx=(8, 0), pady=6)

        ttk.Label(container, text="Excel 文件").grid(row=4, column=0, sticky=tk.W, pady=6)
        ttk.Entry(container, textvariable=self.excel_path_var).grid(row=4, column=1, sticky=tk.EW, pady=6)
        ttk.Button(container, text="选择", command=self.choose_excel).grid(row=4, column=2, padx=(8, 0), pady=6)

        self.start_button = ttk.Button(container, text="开始解析", command=self.start_process)
        self.start_button.grid(row=5, column=1, sticky=tk.W, pady=(12, 6))

        ttk.Label(container, textvariable=self.status_var).grid(row=6, column=0, columnspan=3, sticky=tk.W, pady=6)

        log_frame = ttk.LabelFrame(container, text="运行日志", padding=8)
        log_frame.grid(row=7, column=0, columnspan=3, sticky=tk.NSEW, pady=(10, 0))
        container.rowconfigure(7, weight=1)
        log_frame.rowconfigure(0, weight=1)
        log_frame.columnconfigure(0, weight=1)

        self.log_text = tk.Text(log_frame, height=18, wrap=tk.WORD)
        self.log_text.grid(row=0, column=0, sticky=tk.NSEW)
        self.log_text.tag_configure("success", foreground="#008000")
        self.log_text.tag_configure("fail", foreground="#C00000")
        scrollbar = ttk.Scrollbar(log_frame, command=self.log_text.yview)
        scrollbar.grid(row=0, column=1, sticky=tk.NS)
        self.log_text.configure(yscrollcommand=scrollbar.set)

    def choose_dll(self):
        path = filedialog.askopenfilename(
            title="选择 jsonrichlinksdk.dll",
            filetypes=[("DLL 文件", "*.dll"), ("所有文件", "*.*")],
        )
        if path:
            self.dll_path_var.set(path)

    def choose_excel(self):
        path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            self.excel_path_var.set(path)

    def log(self, message):
        self.after(0, self._append_log, message)

    def _append_log(self, message):
        tag = None
        if "失败" in message or "fail" in message.lower() or "error" in message.lower():
            tag = "fail"
        elif "成功" in message or "success" in message.lower() or "PASS" in message:
            tag = "success"

        if tag:
            self.log_text.insert(tk.END, message + "\n", tag)
        else:
            self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)

    def set_status(self, message):
        self.after(0, self.status_var.set, message)

    def start_process(self):
        if self.worker and self.worker.is_alive():
            return

        try:
            dll_path, excel_path, mcukey, protocol_version, conversion_mode = self.validate_inputs()
        except ValueError as exc:
            messagebox.showerror("参数错误", str(exc))
            return

        self.start_button.configure(state=tk.DISABLED)
        self.log_text.delete("1.0", tk.END)
        self.status_var.set("正在解析，请稍候...")
        self.worker = threading.Thread(
            target=self.run_process,
            args=(dll_path, excel_path, mcukey, protocol_version, conversion_mode),
            daemon=True,
        )
        self.worker.start()

    def validate_inputs(self):
        dll_path = Path(self.dll_path_var.get().strip())
        excel_path = Path(self.excel_path_var.get().strip())
        mcukey = self.mcukey_var.get().strip()

        if not dll_path.is_file():
            raise ValueError("请选择有效的 DLL 文件。")
        if not excel_path.is_file():
            raise ValueError("请选择有效的 Excel 文件。")
        if excel_path.suffix.lower() != ".xlsx":
            raise ValueError("当前只支持 .xlsx 文件。")
        try:
            key_to_bytes(mcukey)
        except ValueError:
            raise ValueError("MCUKEY 必须是有效的 HEX 字符串。")

        protocol_version = int(self.protocol_var.get().split(":", 1)[0])
        conversion_mode = self.mode_var.get()
        return dll_path, excel_path, mcukey, protocol_version, conversion_mode

    def run_process(self, dll_path, excel_path, mcukey, protocol_version, conversion_mode):
        try:
            output_path = build_output_path(excel_path, conversion_mode)
            self.log(f"DLL: {dll_path}")
            self.log(f"Excel: {excel_path}")
            self.log(f"输出: {output_path}")
            self.log(f"转换方向: {conversion_mode}")
            self.log(f"协议: {protocol_version} ({PROTOCOL_VERSIONS[protocol_version]})")

            stats = process_excel(
                dll_path=dll_path,
                excel_path=excel_path,
                output_path=output_path,
                mcukey=mcukey,
                protocol_version=protocol_version,
                conversion_mode=conversion_mode,
                log=self.log,
            )

            self.output_path = output_path
            self.set_status(f"解析完成：{output_path}")
            self.after(
                0,
                lambda: messagebox.showinfo(
                    "完成",
                    f"解析完成。\n\n输出文件：{output_path}\n\n"
                    f"A->B/C 成功：{stats['first_parse_success']}，失败：{stats['first_parse_fail']}\n"
                    f"C->F 成功：{stats['second_parse_success']}，失败：{stats['second_parse_fail']}",
                ),
            )
        except Exception as exc:
            self.set_status("解析失败")
            self.log(f"解析失败: {exc}")
            self.after(0, lambda: messagebox.showerror("解析失败", str(exc)))
        finally:
            self.after(0, lambda: self.start_button.configure(state=tk.NORMAL))


def build_output_path(excel_path, conversion_mode):
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    mode_suffix = "JSON-HEX" if conversion_mode == MODE_JSON_TO_HEX else "HEX-JSON"
    return excel_path.with_name(f"{excel_path.stem}-{mode_suffix}-{timestamp}{excel_path.suffix}")


def parse_hex_to_json(json_sdk, hex_text, mcukey):
    normalized_hex = normalize_hex(hex_text)
    in_buffer = ctypes.create_string_buffer(bytes.fromhex(normalized_hex))
    in_length = len(normalized_hex) // 2
    out_buffer = ctypes.create_string_buffer(OUT_BUFFER_SIZE)
    ctypes.memset(out_buffer, 0, OUT_BUFFER_SIZE)

    parse_result, parsed_data = json_sdk.frame_parse(
        in_buffer,
        in_length,
        out_buffer,
        OUT_BUFFER_SIZE,
        mcukey,
    )
    if parse_result != 0:
        raise RuntimeError(f"{parse_result}: {json_sdk.get_error_description(parse_result)}")

    parsed_json = json.loads(parsed_data)
    return json.dumps(parsed_json, indent=4)


def package_json_to_hex(json_sdk, json_text, mcukey):
    if json_text is None:
        raise ValueError("JSON 内容为空")
    json_text = str(json_text)
    out_buffer = ctypes.create_string_buffer(OUT_BUFFER_SIZE)
    ctypes.memset(out_buffer, 0, OUT_BUFFER_SIZE)
    package_result, packaged_data = json_sdk.frame_package(json_text, out_buffer, OUT_BUFFER_SIZE, mcukey)
    if package_result != 0:
        raise RuntimeError(f"{package_result}: {json_sdk.get_error_description(package_result)}")
    return packaged_data.hex()


def normalize_json_text(value):
    if value is None:
        return ""
    parsed_json = json.loads(str(value))
    return json.dumps(parsed_json, indent=4)


def get_common_resources(dll_path, excel_path, output_path, mcukey, protocol_version, log):
    log_dir = output_path.parent / "logs"
    json_sdk = JsonSdk(dll_path, protocol_version)
    version, code0_desc = json_sdk.initialise(log_dir)
    log(f"SDK Version: {version}")
    log(f"Error Description (Code 0): {code0_desc}")

    workbook = openpyxl.load_workbook(excel_path)
    sheet = workbook.active

    protocol_description = PROTOCOL_VERSIONS.get(protocol_version, "未知协议")
    now_ms = int(datetime.now(timezone.utc).timestamp() * 1000)
    bjtime = (datetime.utcfromtimestamp(now_ms / 1000) + timedelta(hours=8)).strftime("%Y-%m-%d %H:%M:%S")

    sheet.cell(row=1, column=10, value="MCUKEY：" + mcukey)
    sheet.cell(row=2, column=10, value="执行时间：" + bjtime)
    sheet.cell(row=3, column=10, value="协议版本号：" + protocol_description)

    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    green_fill = PatternFill(start_color="FF00B050", end_color="FF00B050", fill_type="solid")
    return json_sdk, workbook, sheet, red_fill, green_fill


def make_stats():
    return {
        "first_parse_success": 0,
        "first_parse_fail": 0,
        "second_parse_success": 0,
        "second_parse_fail": 0,
    }


def fill_result(cell, result, red_fill, green_fill):
    cell.value = result
    cell.fill = green_fill if result == "PASS" else red_fill


def process_excel(dll_path, excel_path, output_path, mcukey, protocol_version, conversion_mode, log):
    if conversion_mode == MODE_JSON_TO_HEX:
        return process_json_to_hex(dll_path, excel_path, output_path, mcukey, protocol_version, log)
    return process_hex_to_json(dll_path, excel_path, output_path, mcukey, protocol_version, log)


def process_hex_to_json(dll_path, excel_path, output_path, mcukey, protocol_version, log):
    json_sdk, workbook, sheet, red_fill, green_fill = get_common_resources(
        dll_path,
        excel_path,
        output_path,
        mcukey,
        protocol_version,
        log,
    )
    stats = make_stats()

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=1, values_only=True), start=1):
        in_buffer_hex = row[0]
        if in_buffer_hex is None:
            continue

        try:
            formatted_json = parse_hex_to_json(json_sdk, in_buffer_hex, mcukey)
            sheet.cell(row=row_idx, column=2, value=formatted_json)
            formatted_hex = package_json_to_hex(json_sdk, formatted_json, mcukey)
            sheet.cell(row=row_idx, column=3, value=formatted_hex)
            stats["first_parse_success"] += 1
            log(f"第 {row_idx} 行 A->B/C 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=2, value=f"data fail: {exc}")
            stats["first_parse_fail"] += 1
            log(f"第 {row_idx} 行 A->B/C 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=3, values_only=True), start=1):
        a_col_value = normalize_hex(row[0])
        c_col_value = normalize_hex(row[2])
        result = "PASS" if c_col_value and a_col_value == c_col_value else "FAIL"
        fill_result(sheet.cell(row=row_idx, column=4), result, red_fill, green_fill)

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, min_col=3, max_col=3, values_only=True), start=1):
        in_buffer_hex = row[0]
        if in_buffer_hex is None:
            continue

        try:
            formatted_json = parse_hex_to_json(json_sdk, in_buffer_hex, mcukey)
            sheet.cell(row=row_idx, column=6, value=formatted_json)

            parsed_json = json.loads(formatted_json)
            msg_id_desc = parsed_json.get("body", {}).get("MsgIDDesc", "")
            sheet.cell(row=row_idx, column=8, value=msg_id_desc)
            stats["second_parse_success"] += 1
            log(f"第 {row_idx} 行 C->F 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=6, value=f"data fail: {exc}")
            stats["second_parse_fail"] += 1
            log(f"第 {row_idx} 行 C->F 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=6, values_only=True), start=1):
        b_col_value = row[1]
        f_col_value = row[5]
        result = "PASS" if b_col_value and b_col_value == f_col_value else "FAIL"
        fill_result(sheet.cell(row=row_idx, column=7), result, red_fill, green_fill)

    workbook.save(output_path)
    return stats


def process_json_to_hex(dll_path, excel_path, output_path, mcukey, protocol_version, log):
    json_sdk, workbook, sheet, red_fill, green_fill = get_common_resources(
        dll_path,
        excel_path,
        output_path,
        mcukey,
        protocol_version,
        log,
    )
    stats = make_stats()

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=1, values_only=True), start=1):
        in_json_data = row[0]
        if in_json_data is None:
            continue

        try:
            formatted_hex = package_json_to_hex(json_sdk, in_json_data, mcukey)
            sheet.cell(row=row_idx, column=2, value=formatted_hex)

            formatted_json = parse_hex_to_json(json_sdk, formatted_hex, mcukey)
            sheet.cell(row=row_idx, column=3, value=formatted_json)

            parsed_json = json.loads(formatted_json)
            msg_id_desc = parsed_json.get("body", {}).get("MsgIDDesc", "")
            sheet.cell(row=row_idx, column=8, value=msg_id_desc)

            stats["first_parse_success"] += 1
            log(f"第 {row_idx} 行 A->B/C 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=2, value=f"data fail: {exc}")
            stats["first_parse_fail"] += 1
            log(f"第 {row_idx} 行 A->B/C 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=3, values_only=True), start=1):
        a_col_value = row[0]
        c_col_value = row[2]
        try:
            result = "PASS" if normalize_json_text(a_col_value) == normalize_json_text(c_col_value) else "FAIL"
        except Exception:
            result = "FAIL"
        fill_result(sheet.cell(row=row_idx, column=4), result, red_fill, green_fill)

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, min_col=3, max_col=3, values_only=True), start=1):
        in_json_data = row[0]
        if in_json_data is None:
            continue

        try:
            formatted_hex = package_json_to_hex(json_sdk, in_json_data, mcukey)
            sheet.cell(row=row_idx, column=6, value=formatted_hex)
            stats["second_parse_success"] += 1
            log(f"第 {row_idx} 行 C->F 成功")
        except Exception as exc:
            sheet.cell(row=row_idx, column=6, value=f"data fail: {exc}")
            stats["second_parse_fail"] += 1
            log(f"第 {row_idx} 行 C->F 失败: {exc}")

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=6, values_only=True), start=1):
        b_col_value = normalize_hex(row[1])
        f_col_value = normalize_hex(row[5])
        result = "PASS" if b_col_value and b_col_value == f_col_value else "FAIL"
        fill_result(sheet.cell(row=row_idx, column=7), result, red_fill, green_fill)

    workbook.save(output_path)
    return stats


def check_user_permission():
    allowed_domain = "RICHLINK.GOSUNCN.COM"
    current_domain = os.environ.get("USERDNSDOMAIN", "").upper()
    if current_domain != allowed_domain:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("提示", "当前用户不在 richlink.gosuncn.com 域内，无法使用")
        root.destroy()
        return False
    return True


if __name__ == "__main__":
    if check_user_permission():
        app = HexJsonGui()
        app.mainloop()
