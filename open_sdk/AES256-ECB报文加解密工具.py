# -*- coding: utf-8 -*-
import os
import re
import threading
import zlib
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl
from openpyxl.styles import PatternFill
from cryptography.hazmat.backends import default_backend
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes


ALLOWED_DOMAIN = "RICHLINK.GOSUNCN.COM"

MODE_DECRYPT = "数据解密"
MODE_ENCRYPT = "数据加密"
MODE_DECRYPT_ENCRYPT = "数据解密后加密"
MODE_CHECK_PAD = "数据校验补位"

FRAME_HEADER_BYTES = 2
PLAIN_HEADER_BYTES = 24
FRAME_END_BYTES = 2
AES_BLOCK_SIZE = 16
# BIT8/9 按当前报文字节展示习惯表示第 8、9 字节，转为 00 00。
BIT89_ZERO_START_BYTE = 7
BIT89_ZERO_BYTE_LEN = 2


def check_user_permission():
    current_domain = os.environ.get("USERDNSDOMAIN", "").upper()
    if current_domain != ALLOWED_DOMAIN:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("提示", "当前用户不在 richlink.gosuncn.com 域内，无法使用")
        root.destroy()
        return False
    return True


def compact_hex(text):
    value = re.sub(r"\s+", "", str(text or ""))
    if len(value) % 2 != 0:
        raise ValueError("HEX 字符串长度必须是偶数。")
    if value and not re.fullmatch(r"[0-9a-fA-F]+", value):
        raise ValueError("数据只能包含 HEX 字符和空白分隔符。")
    return value.lower()


def format_hex(hex_text):
    value = compact_hex(hex_text)
    return value.upper()


def parse_key(key_text, label):
    key_hex = compact_hex(key_text)
    if len(key_hex) != 64:
        raise ValueError(f"{label} 必须是 32 字节 AES-256 KEY，也就是 64 个 HEX 字符。")
    return bytes.fromhex(key_hex)


def split_frame(hex_text, require_aes_block=True):
    value = compact_hex(hex_text)
    min_chars = (PLAIN_HEADER_BYTES + FRAME_END_BYTES) * 2
    if len(value) <= min_chars:
        raise ValueError("数据长度不足，无法按 [48:-4] 取加解密数据。")
    prefix = value[: PLAIN_HEADER_BYTES * 2]
    middle = value[PLAIN_HEADER_BYTES * 2 : -FRAME_END_BYTES * 2]
    suffix = value[-FRAME_END_BYTES * 2 :]
    if require_aes_block and len(middle) // 2 % AES_BLOCK_SIZE != 0:
        raise ValueError("参与 AES 的 [48:-4] 数据长度必须是 16 字节的整数倍。")
    return prefix, middle, suffix


def aes_crypt(hex_text, key_bytes, encrypt=True):
    data = bytes.fromhex(hex_text)
    cipher = Cipher(algorithms.AES(key_bytes), modes.ECB(), backend=default_backend())
    cryptor = cipher.encryptor() if encrypt else cipher.decryptor()
    return (cryptor.update(data) + cryptor.finalize()).hex()


def middle_needs_padding(hex_text):
    _prefix, middle, _suffix = split_frame(hex_text, require_aes_block=False)
    return (len(middle) // 2) % AES_BLOCK_SIZE != 0


def ensure_checksum_padding_frame(hex_text):
    value = compact_hex(hex_text)
    if middle_needs_padding(value):
        return rebuild_checksum_padding_frame(value), True
    return value, False


def rebuild_checksum_padding_frame_with_info(hex_text):
    value = compact_hex(hex_text)
    padding_triggered = middle_needs_padding(value)
    return rebuild_checksum_padding_frame(value), padding_triggered


def force_bit89_zero_frame(hex_text):
    value = compact_hex(hex_text)
    start = BIT89_ZERO_START_BYTE * 2
    end = start + BIT89_ZERO_BYTE_LEN * 2
    if len(value) < end:
        raise ValueError("数据长度不足，无法将 BIT8/9 强转 0000。")
    old_value = value[start:end]
    new_value = value[:start] + "00" * BIT89_ZERO_BYTE_LEN + value[end:]
    return new_value, old_value != "00" * BIT89_ZERO_BYTE_LEN


def process_plain_frame_for_zero_and_checksum(hex_text, force_bit89_zero=False, rebuild_when_forced=True):
    value = compact_hex(hex_text)
    zero_triggered = False
    if force_bit89_zero:
        value, zero_triggered = force_bit89_zero_frame(value)
        if rebuild_when_forced:
            value, padding_triggered = rebuild_checksum_padding_frame_with_info(value)
            return value, padding_triggered, zero_triggered
    value, padding_triggered = ensure_checksum_padding_frame(value)
    return value, padding_triggered, zero_triggered


def build_process_note(padding_triggered, zero_triggered):
    notes = []
    if padding_triggered:
        notes.append("补位")
    if zero_triggered:
        notes.append("触发00")
    return "，".join(notes)


def convert_frame(hex_text, mode, key1_text, key2_text="", force_bit89_zero=False):
    result, _padding_triggered, _zero_triggered = convert_frame_with_info(
        hex_text,
        mode,
        key1_text,
        key2_text,
        force_bit89_zero,
    )
    return result


def convert_frame_with_info(hex_text, mode, key1_text, key2_text="", force_bit89_zero=False):
    if mode == MODE_CHECK_PAD:
        value = compact_hex(hex_text)
        zero_triggered = False
        if force_bit89_zero:
            value, zero_triggered = force_bit89_zero_frame(value)
        rebuilt, padding_triggered = rebuild_checksum_padding_frame_with_info(value)
        return rebuilt, padding_triggered, zero_triggered

    padding_triggered = False
    zero_triggered = False

    if mode == MODE_DECRYPT:
        key1 = parse_key(key1_text, "KEY")
        prefix, middle, suffix = split_frame(hex_text)
        converted_middle = aes_crypt(middle, key1, encrypt=False)
        converted_frame = prefix + converted_middle + suffix
        converted_frame, padding_triggered, zero_triggered = process_plain_frame_for_zero_and_checksum(
            converted_frame,
            force_bit89_zero,
        )
    elif mode == MODE_ENCRYPT:
        key1 = parse_key(key1_text, "KEY")
        padded_frame, padding_triggered, zero_triggered = process_plain_frame_for_zero_and_checksum(
            hex_text,
            force_bit89_zero,
        )
        prefix, middle, suffix = split_frame(padded_frame)
        converted_middle = aes_crypt(middle, key1, encrypt=True)
        converted_frame = prefix + converted_middle + suffix
    elif mode == MODE_DECRYPT_ENCRYPT:
        key1 = parse_key(key1_text, "KEY1")
        key2 = parse_key(key2_text, "KEY2")
        prefix, middle, suffix = split_frame(hex_text)
        decrypted_middle = aes_crypt(middle, key1, encrypt=False)
        decrypted_frame = prefix + decrypted_middle + suffix
        decrypted_frame, padding_triggered, zero_triggered = process_plain_frame_for_zero_and_checksum(
            decrypted_frame,
            force_bit89_zero,
        )
        prefix, middle, suffix = split_frame(decrypted_frame)
        converted_middle = aes_crypt(middle, key2, encrypt=True)
        converted_frame = prefix + converted_middle + suffix
    else:
        raise ValueError("未知转换模式。")

    return converted_frame, padding_triggered, zero_triggered


def find_checksum_bounds(hex_text):
    value = compact_hex(hex_text)
    try:
        frame_bytes = bytes.fromhex(value)
        effective_len = int.from_bytes(frame_bytes[27:29], "big")
        checksum_start_byte = 29 + effective_len
        checksum_end_byte = checksum_start_byte + 4
        frame_end_byte = len(frame_bytes) - FRAME_END_BYTES
        if PLAIN_HEADER_BYTES <= checksum_start_byte and checksum_end_byte <= frame_end_byte:
            return checksum_start_byte * 2, checksum_end_byte * 2
    except Exception:
        pass
    return len(value) - 12, len(value) - 4


def rebuild_checksum_padding_frame(hex_text):
    value = compact_hex(hex_text)
    if not value.startswith("5555"):
        raise ValueError("数据必须以 5555 开头。")
    if not value.endswith("aaaa"):
        raise ValueError("数据必须以 aaaa 结尾。")
    if len(value) < 60:
        raise ValueError("数据长度不足，无法提取头部、校验位和帧尾。")

    checksum_start, checksum_end = find_checksum_bounds(value)
    unencrypted_data = value[PLAIN_HEADER_BYTES * 2 : checksum_start]
    if checksum_end > len(value) - FRAME_END_BYTES * 2:
        raise ValueError("校验位位置异常，无法重算。")

    old_checksum = value[checksum_start:checksum_end]
    plaintext_with_checksum = unencrypted_data + old_checksum
    plaintext_len = len(plaintext_with_checksum) // 2
    padding_len = (AES_BLOCK_SIZE - plaintext_len % AES_BLOCK_SIZE) % AES_BLOCK_SIZE
    supplement = "ff" * padding_len

    frame_without_new_checksum = value[: PLAIN_HEADER_BYTES * 2] + plaintext_with_checksum + supplement + value[-4:]
    message_length = len(frame_without_new_checksum) // 2
    message_length_hex = f"{message_length:04x}"

    header_data_new = (
        message_length_hex
        + value[8 : PLAIN_HEADER_BYTES * 2]
    )
    checksum_source = header_data_new + unencrypted_data
    checksum = zlib.adler32(bytes.fromhex(checksum_source))
    checksum_hex = f"{checksum:08x}"

    return value[:4] + header_data_new + unencrypted_data + checksum_hex + supplement + value[-4:]


def build_output_path(excel_path, mode):
    suffix_map = {
        MODE_DECRYPT: "DECRYPT",
        MODE_ENCRYPT: "ENCRYPT",
        MODE_DECRYPT_ENCRYPT: "DECRYPT-ENCRYPT",
        MODE_CHECK_PAD: "CHECK-PAD",
    }
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    return excel_path.with_name(f"{excel_path.stem}-{suffix_map[mode]}-{timestamp}{excel_path.suffix}")


def byte_ranges_for_coloring(hex_text, encrypted_payload=False):
    value = compact_hex(hex_text)
    total_bytes = len(value) // 2
    ranges = []

    def add(start, end, tag):
        if start < end and start < total_bytes:
            ranges.append((start, min(end, total_bytes), tag))

    add(0, 2, "frame")
    add(2, 4, "length")
    add(4, 6, "protocol")
    add(6, 7, "encrypt_type")
    add(7, 8, "diagnostic")
    add(8, 9, "battery")
    add(9, 24, "device")

    body_start = 24
    frame_end_start = max(total_bytes - 2, 0)
    add(frame_end_start, total_bytes, "frame")

    if total_bytes <= body_start:
        return ranges

    if encrypted_payload:
        add(body_start, frame_end_start, "cipher")
        return ranges

    add(24, 25, "frame_type")
    add(25, 27, "frame_number")
    add(27, 29, "effective_length")

    try:
        bytes_data = bytes.fromhex(value)
        effective_len = int.from_bytes(bytes_data[27:29], "big")
        valid_start = 29
        valid_end = valid_start + effective_len
        checksum_end = valid_end + 4
        if valid_end <= frame_end_start:
            add(valid_start, valid_end, "valid")
            add(valid_end, min(checksum_end, frame_end_start), "checksum")
            add(min(checksum_end, frame_end_start), frame_end_start, "padding")
        else:
            add(body_start, frame_end_start, "payload")
    except Exception:
        add(body_start, frame_end_start, "payload")

    return ranges


class AesFrameTool(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("AES-256-ECB 报文加解密工具-Mangetic-2026-8-14-域控")
        self.geometry("960x720")
        self.minsize(880, 640)

        self.mode_var = tk.StringVar(value=MODE_DECRYPT)
        self.key1_var = tk.StringVar()
        self.key2_var = tk.StringVar()
        self.force_zero_var = tk.BooleanVar(value=False)
        self.excel_path_var = tk.StringVar()
        self.status_var = tk.StringVar(value="请输入数据和 KEY，或选择 Excel 批量转换")
        self.warning_var = tk.StringVar()
        self.worker = None
        self.color_tags = (
            "frame",
            "length",
            "protocol",
            "encrypt_type",
            "diagnostic",
            "battery",
            "device",
            "frame_type",
            "frame_number",
            "effective_length",
            "valid",
            "checksum",
            "padding",
            "cipher",
            "payload",
        )

        self.create_widgets()
        self.on_mode_changed()

    def create_widgets(self):
        container = ttk.Frame(self, padding=14)
        container.pack(fill=tk.BOTH, expand=True)
        container.columnconfigure(1, weight=1)
        container.rowconfigure(7, weight=1)
        container.rowconfigure(9, weight=1)

        ttk.Label(container, text="转换模式").grid(row=0, column=0, sticky=tk.W, pady=5)
        mode_frame = ttk.Frame(container)
        mode_frame.grid(row=0, column=1, columnspan=2, sticky=tk.W, pady=5)
        mode_box = ttk.Combobox(
            mode_frame,
            textvariable=self.mode_var,
            values=[MODE_DECRYPT, MODE_ENCRYPT, MODE_DECRYPT_ENCRYPT, MODE_CHECK_PAD],
            state="readonly",
            width=24,
        )
        mode_box.pack(side=tk.LEFT)
        mode_box.bind("<<ComboboxSelected>>", lambda _event: self.on_mode_changed(clear_text=True))
        ttk.Checkbutton(
            mode_frame,
            text="BIT8/9强转0000",
            variable=self.force_zero_var,
        ).pack(side=tk.LEFT, padx=(8, 0))

        self.key1_label = ttk.Label(container, text="KEY")
        self.key1_entry = ttk.Entry(container, textvariable=self.key1_var)
        self.key1_label.grid(row=1, column=0, sticky=tk.W, pady=5)
        self.key1_entry.grid(row=1, column=1, columnspan=2, sticky=tk.EW, pady=5)

        self.key2_label = ttk.Label(container, text="KEY2(加密)")
        self.key2_entry = ttk.Entry(container, textvariable=self.key2_var)
        self.key2_label.grid(row=2, column=0, sticky=tk.W, pady=5)
        self.key2_entry.grid(row=2, column=1, columnspan=2, sticky=tk.EW, pady=5)

        ttk.Label(container, text="Excel 文件").grid(row=3, column=0, sticky=tk.W, pady=5)
        ttk.Entry(container, textvariable=self.excel_path_var).grid(row=3, column=1, sticky=tk.EW, pady=5)
        ttk.Button(container, text="选择", command=self.choose_excel).grid(row=3, column=2, padx=(8, 0), pady=5)

        button_frame = ttk.Frame(container)
        button_frame.grid(row=4, column=1, sticky=tk.W, pady=(8, 6))
        ttk.Button(button_frame, text="转换文本框数据", command=self.convert_text_data).pack(side=tk.LEFT)
        ttk.Button(button_frame, text="批量转换 Excel", command=self.start_batch_convert).pack(side=tk.LEFT, padx=(10, 0))

        ttk.Label(container, text="输入数据").grid(row=6, column=0, sticky=tk.W, pady=(10, 4))
        self.input_text = tk.Text(container, height=11, wrap=tk.WORD)
        self.input_text.grid(row=7, column=0, columnspan=3, sticky=tk.NSEW)

        ttk.Label(container, text="输出数据").grid(row=8, column=0, sticky=tk.W, pady=(10, 4))
        self.output_text = tk.Text(container, height=11, wrap=tk.WORD)
        self.output_text.grid(row=9, column=0, columnspan=3, sticky=tk.NSEW)

        ttk.Label(container, textvariable=self.status_var).grid(row=10, column=0, columnspan=3, sticky=tk.W, pady=(10, 0))
        tk.Label(
            container,
            textvariable=self.warning_var,
            fg="#C00000",
            font=("Microsoft YaHei UI", 10, "bold"),
        ).grid(row=11, column=0, columnspan=3, sticky=tk.W, pady=(4, 0))

        for text_widget in (self.input_text, self.output_text):
            self.configure_text_tags(text_widget)
            text_widget.bind("<<Paste>>", lambda event, widget=text_widget: self.after(80, self.format_and_color, widget))
            text_widget.bind("<KeyRelease>", lambda event, widget=text_widget: self.after(80, self.color_text, widget))

    def configure_text_tags(self, text_widget):
        text_widget.tag_configure("frame", foreground="#C00000")
        text_widget.tag_configure("length", foreground="#0070C0")
        text_widget.tag_configure("protocol", foreground="#7030A0")
        text_widget.tag_configure("encrypt_type", foreground="#9E480E")
        text_widget.tag_configure("diagnostic", foreground="#8064A2")
        text_widget.tag_configure("battery", foreground="#00A2E8")
        text_widget.tag_configure("device", foreground="#00B050")
        text_widget.tag_configure("frame_type", foreground="#C55A11")
        text_widget.tag_configure("frame_number", foreground="#2F5597")
        text_widget.tag_configure("effective_length", foreground="#E36C0A")
        text_widget.tag_configure("valid", foreground="#000000")
        text_widget.tag_configure("checksum", foreground="#FF00FF")
        text_widget.tag_configure("padding", foreground="#808080")
        text_widget.tag_configure("cipher", foreground="#7F6000")
        text_widget.tag_configure("payload", foreground="#666666")

    def on_mode_changed(self, clear_text=False):
        mode = self.mode_var.get()
        if mode == MODE_DECRYPT_ENCRYPT:
            self.key1_label.configure(text="KEY1（解密）")
            self.key2_label.configure(text="KEY2(加密)")
            self.key1_label.grid()
            self.key1_entry.grid()
            self.key2_label.grid()
            self.key2_entry.grid()
        elif mode == MODE_CHECK_PAD:
            self.key1_label.grid_remove()
            self.key1_entry.grid_remove()
            self.key2_label.grid_remove()
            self.key2_entry.grid_remove()
        else:
            self.key1_label.configure(text="KEY")
            self.key1_label.grid()
            self.key1_entry.grid()
            self.key2_label.grid_remove()
            self.key2_entry.grid_remove()
        if clear_text:
            self.input_text.delete("1.0", tk.END)
            self.output_text.delete("1.0", tk.END)
            self.warning_var.set("")
            self.status_var.set("已切换模式，请重新输入待转换数据")

    def choose_excel(self):
        path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            self.excel_path_var.set(path)

    def text_value(self, text_widget):
        return text_widget.get("1.0", tk.END).strip()

    def set_text_value(self, text_widget, value):
        text_widget.delete("1.0", tk.END)
        text_widget.insert("1.0", format_hex(value))
        self.color_text(text_widget, force_plain=self.mode_var.get() in (MODE_DECRYPT, MODE_CHECK_PAD))

    def format_and_color(self, text_widget):
        try:
            value = compact_hex(self.text_value(text_widget))
            if value:
                text_widget.delete("1.0", tk.END)
                text_widget.insert("1.0", format_hex(value))
        except Exception:
            pass
        self.color_text(text_widget)

    def color_text(self, text_widget, force_plain=None):
        for tag in self.color_tags:
            text_widget.tag_remove(tag, "1.0", tk.END)

        raw_text = self.text_value(text_widget)
        try:
            formatted = format_hex(raw_text)
        except Exception:
            return

        if formatted != raw_text and "\n" not in raw_text:
            return

        encrypted_payload = False
        if force_plain is None:
            encrypted_payload = self.mode_var.get() == MODE_ENCRYPT and text_widget is self.output_text
        else:
            encrypted_payload = not force_plain

        byte_ranges = byte_ranges_for_coloring(formatted, encrypted_payload=encrypted_payload)
        for start_byte, end_byte, tag in byte_ranges:
            start_index = self.byte_to_text_index(start_byte)
            end_index = self.byte_to_text_index(end_byte)
            text_widget.tag_add(tag, start_index, end_index)

    @staticmethod
    def byte_to_text_index(byte_index):
        if byte_index <= 0:
            return "1.0"
        return f"1.{byte_index * 2}"

    def convert_text_data(self):
        try:
            result, padding_triggered, zero_triggered = convert_frame_with_info(
                self.text_value(self.input_text),
                self.mode_var.get(),
                self.key1_var.get(),
                self.key2_var.get(),
                self.force_zero_var.get(),
            )
            self.set_text_value(self.output_text, result)
            self.format_and_color(self.input_text)
            process_note = build_process_note(padding_triggered, zero_triggered)
            if process_note:
                self.warning_var.set(f"处理提示：{process_note}")
            else:
                self.warning_var.set("")
            self.status_var.set("文本框数据转换完成")
        except Exception as exc:
            messagebox.showerror("转换失败", str(exc))
            self.warning_var.set("")
            self.status_var.set("文本框数据转换失败")

    def start_batch_convert(self):
        if self.worker and self.worker.is_alive():
            return
        excel_path = Path(self.excel_path_var.get().strip())
        if not excel_path.is_file():
            messagebox.showerror("参数错误", "请选择有效的 Excel 文件。")
            return
        if excel_path.suffix.lower() != ".xlsx":
            messagebox.showerror("参数错误", "当前只支持 .xlsx 文件。")
            return

        self.worker = threading.Thread(target=self.batch_convert, args=(excel_path,), daemon=True)
        self.worker.start()

    def batch_convert(self, excel_path):
        try:
            output_path = build_output_path(excel_path, self.mode_var.get())
            workbook = openpyxl.load_workbook(excel_path)
            sheet = workbook.active

            success_count = 0
            fail_count = 0
            padding_count = 0
            zero_count = 0
            red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
            for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=1, values_only=True), start=1):
                source_data = row[0]
                if source_data is None:
                    continue
                try:
                    result, padding_triggered, zero_triggered = convert_frame_with_info(
                        source_data,
                        self.mode_var.get(),
                        self.key1_var.get(),
                        self.key2_var.get(),
                        self.force_zero_var.get(),
                    )
                    sheet.cell(row=row_idx, column=2, value=format_hex(result))
                    process_note = build_process_note(padding_triggered, zero_triggered)
                    if process_note:
                        note_cell = sheet.cell(row=row_idx, column=3, value=process_note)
                        note_cell.fill = red_fill
                    if padding_triggered:
                        padding_count += 1
                    if zero_triggered:
                        zero_count += 1
                    success_count += 1
                except Exception as exc:
                    sheet.cell(row=row_idx, column=2, value=f"data fail: {exc}")
                    fail_count += 1

            workbook.save(output_path)
            self.after(
                0,
                lambda: self.status_var.set(
                    f"批量转换完成：成功 {success_count}，失败 {fail_count}，补位 {padding_count}，触发00 {zero_count}"
                ),
            )
            self.after(
                0,
                lambda: messagebox.showinfo(
                    "完成",
                    f"输出文件：{output_path}\n\n"
                    f"解析数据：{success_count + fail_count} 条\n"
                    f"成功：{success_count} 条\n"
                    f"失败：{fail_count} 条\n"
                    f"补位：{padding_count} 条\n"
                    f"触发00：{zero_count} 条",
                ),
            )
        except Exception as exc:
            self.after(0, lambda: self.status_var.set("批量转换失败"))
            self.after(0, lambda: messagebox.showerror("批量转换失败", str(exc)))


if __name__ == "__main__":
    if check_user_permission():
        app = AesFrameTool()
        app.mainloop()


# 规则逻辑说明：
# 1. 报文整体结构：
#    55 55 + Message Length + Protocol Version + Encryption Type + Self Diagnostic
#    + Battery Level + Device ID/IMEI + Frame Type + Frame Number
#    + Effective Data Length + Valid Data + Adler32 + FF补位 + aa aa
#
# 2. Message Length：
#    第 3-4 字节，表示整帧长度，从 55 55 帧头开始，到 aa aa 帧尾结束。
#
# 3. Effective Data Length：
#    第 28-29 字节，只表示 valid_data 的长度。
#    valid_data 不包含 Frame Type、Frame Number、Effective Data Length、Adler32、FF补位、帧头帧尾。
#
# 4. Adler32 校验位：
#    Adler32 不是加密，不可逆，只用于校验数据是否变化。
#    校验通常从 Message Length 开始算，也就是不包含 55 55 帧头；
#    计算内容包含 header_data_new + unencrypted_data：
#    Message Length + Protocol Version + Encryption Type + Self Diagnostic
#    + Battery Level + Device ID/IMEI + Frame Type + Frame Number
#    + Effective Data Length + valid_data。
#    不包含 Adler32 自己、FF补位、aa aa 帧尾。
#
# 5. 补位：
#    AES 要求输入长度是 16 字节整数倍。
#    补位是对即将加密的明文段计算：
#    Frame Type + Frame Number + Effective Data Length + valid_data + Adler32。
#    如果长度不足 16 字节倍数，就在 Adler32 后面、aa aa 前面补 FF。
#
# 6. 加密：
#    使用 AES-256-ECB，KEY 长度必须是 32 字节，也就是 64 个 HEX 字符。
#    加密范围固定为 compact_hex_data[48:-4]：
#    即从第 25 字节开始，到 aa aa 帧尾之前。
#    前 24 字节明文头和最后 2 字节 aa aa 不加密。
#
# 7. 解密：
#    解密范围与加密范围相同，也是 compact_hex_data[48:-4]。
#    使用同一个 AES-256-ECB KEY 可还原这段密文。
#
# 8. 数据校验补位模式：
#    该模式不需要 KEY，适合处理明文报文长度不是 16 字节倍数的情况。
#    先定位原报文中的 4 字节 Adler32 校验位：
#    优先通过第 28-29 字节 Effective Data Length 计算校验位位置；
#    如果无法定位，则按 aa aa 前 4 字节，即 compact_hex_data[-12:-4] 作为旧校验位。
#    然后取第 25 字节开始到旧校验位前的数据：
#    Frame Type + Frame Number + Effective Data Length + valid_data。
#    旧校验位只用于计算补位长度占位，不参与最终校验。
#    将上述数据 + 4 字节校验位占位后的长度补 FF 到 16 字节倍数。
#    再根据补位后的完整帧长度替换第 3-4 字节 Message Length，
#    即 HEX 字符串下标 [4:8]。
#    最后用替换长度后的 header_data_new + unencrypted_data 重算 Adler32：
#    不包含 55 55、旧校验位、FF 补位、aa aa。
#    输出完整数据：
#    55 55 + 新长度/头部 + 明文业务数据 + 新 Adler32 + FF补位 + aa aa。
