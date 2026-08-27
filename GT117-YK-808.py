# -*- coding: utf-8 -*-
"""
JT/T 808 协议 HEX 解析 GUI 工具。

功能重点：
1. 输入的 HEX 会先自动去掉回车、换行、空格、Tab。
2. 支持 7E 帧边界、7D01/7D02 转义还原、异或校验。
3. 解析 JT/T 808-2013 常规消息头，并兼容 2019 版本标识位的头部长度。
4. 对常见标准消息体做结构化解析，未知消息体保留原始 HEX。
"""

import json
import os
import re
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, scrolledtext, ttk


SAMPLE_HEX = "7E00010005115436621650000220D1814000617E"

MESSAGE_ID_NAMES = {
    0x0001: "终端通用应答",
    0x0002: "终端心跳",
    0x0003: "终端注销",
    0x0004: "查询服务器时间",
    0x0005: "终端补传分包请求",
    0x0100: "终端注册",
    0x0102: "终端鉴权",
    0x0104: "查询终端参数应答",
    0x0107: "查询终端属性应答",
    0x0108: "终端升级结果通知",
    0x0200: "位置信息汇报",
    0x0201: "位置信息查询应答",
    0x0303: "信息点播/取消",
    0x0500: "车辆控制应答",
    0x0700: "行驶记录仪数据上传",
    0x0702: "驾驶员身份信息采集上报",
    0x0301: "事件报告",
    0x0302: "提问应答",
    0x0704: "定位数据批量上传",
    0x0705: "CAN总线数据上传",
    0x0800: "多媒体事件信息上传",
    0x0801: "多媒体数据上传",
    0x8001: "平台通用应答",
    0x8003: "服务器补传分包请求",
    0x8004: "查询服务器时间应答",
    0x8100: "终端注册应答",
    0x8103: "设置终端参数",
    0x8104: "查询终端参数",
    0x8105: "终端控制",
    0x8106: "查询指定终端参数",
    0x8107: "查询终端属性",
    0x8108: "下发终端升级包",
    0x8201: "位置信息查询",
    0x8202: "临时位置跟踪控制",
    0x8203: "人工确认报警消息",
    0x8204: "服务器向终端发起链路检测请求",
    0x8300: "文本信息下发",
    0x8301: "事件设置",
    0x8302: "提问下发",
    0x8303: "信息点播菜单设置",
    0x8304: "信息服务",
    0x8400: "电话回拨",
    0x8401: "设置电话本",
    0x8500: "车辆控制",
    0x8600: "设置圆形区域",
    0x8601: "删除圆形区域",
    0x8602: "设置矩形区域",
    0x8603: "删除矩形区域",
    0x8604: "设置多边形区域",
    0x8605: "删除多边形区域",
    0x8606: "设置路线",
    0x8607: "删除路线",
    0x8700: "行驶记录仪数据采集命令",
    0x8701: "行驶记录仪参数下传命令",
    0x8702: "上报驾驶员身份信息请求",
    0x8800: "多媒体数据上传应答",
    0x8801: "摄像头立即拍摄命令",
    0x0805: "摄像头立即拍摄命令应答",
}

RESULT_NAMES = {
    0x00: "Success",
    0x01: "Failure",
    0x02: "Message error",
    0x03: "Unsupported",
    0x04: "Alarm processing confirmation",
}

REGISTER_RESULT_NAMES = {
    0x00: "Success",
    0x01: "Vehicle registered",
    0x02: "No vehicle in database",
    0x03: "Terminal registered",
    0x04: "No terminal in database",
}

ENCRYPTION_NAMES = {
    0: "None",
    1: "RSA",
}

ALARM_BIT_NAMES = {
    0: "紧急报警",
    1: "超速报警",
    2: "疲劳驾驶",
    3: "危险预警",
    4: "GNSS模块故障",
    5: "GNSS天线未接或被剪断",
    6: "GNSS天线短路",
    7: "终端主电源欠压",
    8: "终端主电源掉电",
    9: "终端LCD或显示器故障",
    10: "TTS模块故障",
    11: "摄像头故障",
    18: "当天累计驾驶超时",
    19: "超时停车",
    20: "进出区域",
    21: "进出路线",
    22: "路段行驶时间不足/过长",
    23: "路线偏离报警",
    24: "车辆VSS故障",
    25: "车辆油量异常",
    26: "车辆被盗",
    27: "车辆非法点火",
    28: "车辆非法位移",
    29: "碰撞侧翻报警",
    30: "非法开门报警",
}

STATUS_BIT_NAMES = {
    0: "ACC开",
    1: "定位",
    2: "南纬",
    3: "西经",
    4: "运营状态",
    5: "经纬度已保密插件加密",
    10: "车辆油路断开",
    11: "车辆电路断开",
    12: "车门加锁",
    13: "前门开",
    14: "中门开",
    15: "后门开",
    16: "驾驶席门开",
    17: "自定义门开",
    18: "使用GPS卫星定位",
    19: "使用北斗卫星定位",
    20: "使用GLONASS卫星定位",
    21: "使用Galileo卫星定位",
}

TERMINAL_TYPE_BIT_NAMES = {
    0: "适用客运车辆",
    1: "适用危险品车辆",
    2: "适用普通货运车辆",
    3: "适用出租车辆",
    6: "支持硬盘录像",
    7: "分体机",
    8: "适用挂车",
}

GNSS_MODULE_BIT_NAMES = {
    0: "支持GPS定位",
    1: "支持北斗定位",
    2: "支持GLONASS定位",
    3: "支持Galileo定位",
}

COMM_MODULE_BIT_NAMES = {
    0: "支持GPRS通信",
    1: "支持CDMA通信",
    2: "支持TD-SCDMA通信",
    3: "支持WCDMA通信",
    4: "支持CDMA2000通信",
    5: "支持TD-LTE通信",
    7: "支持其他通信方式",
}

UPGRADE_TYPE_NAMES = {
    0x00: "终端",
    0x0C: "道路运输证IC卡读卡器",
    0x34: "卫星定位模块/北斗卫星定位模块",
}

UPGRADE_RESULT_NAMES = {
    0x00: "成功",
    0x01: "失败",
    0x02: "取消",
}

PHONEBOOK_SETTING_TYPE_NAMES = {
    0x00: "删除终端上所有存储的联系人",
    0x01: "更新电话本",
    0x02: "追加电话本",
    0x03: "修改电话本",
}


class ParseError(ValueError):
    """用于把解析错误清晰显示到 GUI。"""


def clean_hex_text(text):
    """去掉常见空白，并兼容用户从文档里复制出的引号和 0x 前缀。"""
    cleaned = text.strip().strip("\"'“”‘’")
    cleaned = re.sub(r"\s+", "", cleaned)
    cleaned = re.sub(r"(?i)0x", "", cleaned)
    return cleaned.upper()


def ensure_hex_bytes(hex_text):
    cleaned = clean_hex_text(hex_text)
    if not cleaned:
        raise ParseError("请输入 HEX 数据。")
    if len(cleaned) % 2:
        raise ParseError("HEX 长度必须为偶数。")
    if re.search(r"[^0-9A-F]", cleaned):
        raise ParseError("输入中包含非 HEX 字符，请只输入 0-9、A-F。")
    try:
        return cleaned, bytes.fromhex(cleaned)
    except ValueError as exc:
        raise ParseError(f"HEX 转字节失败：{exc}") from exc


def xor_checksum(data):
    value = 0
    for item in data:
        value ^= item
    return value


def unescape_808(data):
    result = bytearray()
    i = 0
    while i < len(data):
        current = data[i]
        if current == 0x7D:
            if i + 1 >= len(data):
                raise ParseError("转义字节 7D 后缺少后续字节。")
            next_byte = data[i + 1]
            if next_byte == 0x02:
                result.append(0x7E)
            elif next_byte == 0x01:
                result.append(0x7D)
            else:
                raise ParseError(f"非法转义序列：7D{next_byte:02X}")
            i += 2
        else:
            result.append(current)
            i += 1
    return bytes(result)


def bytes_to_bcd(data):
    chars = []
    for item in data:
        chars.append(f"{(item >> 4) & 0x0F:X}")
        chars.append(f"{item & 0x0F:X}")
    return "".join(chars)


def read_u8(data, offset):
    require_length(data, offset, 1)
    return data[offset], offset + 1


def read_u16(data, offset):
    require_length(data, offset, 2)
    return int.from_bytes(data[offset:offset + 2], "big"), offset + 2


def read_u32(data, offset):
    require_length(data, offset, 4)
    return int.from_bytes(data[offset:offset + 4], "big"), offset + 4


def read_bytes(data, offset, size):
    require_length(data, offset, size)
    return data[offset:offset + size], offset + size


def require_length(data, offset, size):
    if offset + size > len(data):
        raise ParseError(f"数据长度不足：偏移 {offset} 需要 {size} 字节，实际剩余 {len(data) - offset} 字节。")


def hex_slice(data):
    return data.hex().upper()


def safe_decode(data):
    stripped = data.rstrip(b"\x00")
    if not stripped:
        return ""
    for encoding in ("gbk", "ascii", "utf-8"):
        try:
            return stripped.decode(encoding)
        except UnicodeDecodeError:
            continue
    return hex_slice(data)


def format_bits(value, width):
    return f"{value:0{width}b}"


def parse_message_body_property(prop, protocol="2013"):
    binary = format_bits(prop, 16)
    encrypt_type = (prop >> 10) & 0x07
    version_flag = protocol == "2019" and bool((prop >> 14) & 0x01)
    return {
        f"[{binary}]消息体属性": prop,
        "协议标准": "JTT2019" if protocol == "2019" else "JTT2013",
        "版本号": "JTT2019" if version_flag else "JTT2013",
        "[bit15]保留": (prop >> 15) & 0x01,
        "[bit14]版本标识" if protocol == "2019" else "[bit14]保留": (prop >> 14) & 0x01,
        "[bit13]是否分包": bool((prop >> 13) & 0x01),
        "[bit10~bit12]数据加密": ENCRYPTION_NAMES.get(encrypt_type, f"Unknown({encrypt_type})"),
        "[bit0~bit9]消息体长度": prop & 0x03FF,
    }


def parse_header(packet_without_checksum, protocol="2013"):
    msg_id, offset = read_u16(packet_without_checksum, 0)
    prop, offset = read_u16(packet_without_checksum, offset)
    prop_obj = parse_message_body_property(prop, protocol)
    version_flag = protocol == "2019" and bool((prop >> 14) & 0x01)

    protocol_version = None
    if version_flag:
        protocol_version, offset = read_u8(packet_without_checksum, offset)
        phone_len = 10
    else:
        phone_len = 6

    require_length(packet_without_checksum, offset, phone_len + 2)
    phone_raw = packet_without_checksum[offset:offset + phone_len]
    phone = bytes_to_bcd(phone_raw).lstrip("0") or "0"
    offset += phone_len
    serial, offset = read_u16(packet_without_checksum, offset)

    package_info = None
    if (prop >> 13) & 0x01:
        total, offset = read_u16(packet_without_checksum, offset)
        index, offset = read_u16(packet_without_checksum, offset)
        package_info = {
            f"[{total:04X}]消息包总数": total,
            f"[{index:04X}]包序号": index,
        }

    body_length = prop & 0x03FF
    body = packet_without_checksum[offset:offset + body_length]
    if len(body) != body_length:
        raise ParseError(f"消息体长度不足：属性中声明 {body_length} 字节，实际 {len(body)} 字节。")

    extra = packet_without_checksum[offset + body_length:]
    if extra:
        raise ParseError(f"消息体后存在多余数据：{hex_slice(extra)}")

    header = {
        f"[{msg_id:04X}]消息Id": msg_id,
        "消息名称": MESSAGE_ID_NAMES.get(msg_id, f"Unknown(0x{msg_id:04X})"),
        "解析标准": "JT/T 808-2019" if protocol == "2019" else "JT/T 808-2013",
        "消息体属性对象": prop_obj,
    }
    if protocol_version is not None:
        header[f"[{protocol_version:02X}]协议版本号"] = protocol_version
    header[f"[{hex_slice(phone_raw)}]终端手机号"] = phone
    header[f"[{serial:04X}]消息流水号"] = serial
    if package_info:
        header["分包对象"] = package_info

    return msg_id, body, header


def parse_common_response(body, response_name):
    if len(body) != 5:
        return {
            response_name: hex_slice(body),
            "解析提示": f"标准通用应答长度应为 5 字节，当前为 {len(body)} 字节。",
        }
    response_serial = int.from_bytes(body[0:2], "big")
    response_msg_id = int.from_bytes(body[2:4], "big")
    result = body[4]
    return {
        response_name: hex_slice(body),
        f"[{body[0:2].hex().upper()}]应答流水号": response_serial,
        f"[{body[2:4].hex().upper()}]应答消息Id": response_msg_id,
        f"[{body[4]:02X}]结果": RESULT_NAMES.get(result, f"Unknown({result})"),
    }


def parse_terminal_register(body, protocol="2013"):
    result = {"终端注册": hex_slice(body)}
    if protocol == "2019":
        manufacturer_len = 11
        model_len = 30
        terminal_id_len = 30
        min_len = 76
        layout_name = "JT/T 808-2019：制造商ID11字节，型号30字节，终端ID30字节"
    else:
        manufacturer_len = 5
        model_len = 20
        terminal_id_len = 7
        min_len = 37
        layout_name = "JT/T 808-2013：制造商ID5字节，型号20字节，终端ID7字节"

    if len(body) < min_len:
        result["解析提示"] = f"{layout_name}，最短 {min_len} 字节，当前为 {len(body)} 字节。"
        return result

    offset = 0
    province, offset = read_u16(body, offset)
    city, offset = read_u16(body, offset)
    manufacturer, offset = read_bytes(body, offset, manufacturer_len)
    model, offset = read_bytes(body, offset, model_len)
    terminal_id, offset = read_bytes(body, offset, terminal_id_len)
    plate_color, offset = read_u8(body, offset)
    plate = body[offset:]

    result.update({
        "字段长度布局": layout_name,
        f"[{body[0:2].hex().upper()}]省域ID": province,
        f"[{body[2:4].hex().upper()}]市县域ID": city,
        f"[{hex_slice(manufacturer)}]制造商ID": safe_decode(manufacturer),
        f"[{hex_slice(model)}]终端型号": safe_decode(model),
        f"[{hex_slice(terminal_id)}]终端ID": safe_decode(terminal_id),
        f"[{plate_color:02X}]车牌颜色": plate_color,
        f"[{hex_slice(plate)}]车辆标识": safe_decode(plate),
    })
    return result


def parse_authentication(body, protocol="2013"):
    if protocol != "2019":
        return {
            "终端鉴权": hex_slice(body),
            "字段长度布局": "JT/T 808-2013：消息体直接为鉴权码",
            f"[{hex_slice(body)}]鉴权码": safe_decode(body),
        }

    result = {
        "终端鉴权": hex_slice(body),
        "字段长度布局": "JT/T 808-2019：鉴权码长度 + 鉴权码 + IMEI + 软件版本号",
    }
    if len(body) < 36:
        result["解析提示"] = f"2019 鉴权消息至少需要 1+n+15+20 字节，当前为 {len(body)} 字节。"
        result[f"[{hex_slice(body)}]原始鉴权数据"] = safe_decode(body)
        return result

    offset = 0
    auth_len, offset = read_u8(body, offset)
    if len(body) < 1 + auth_len + 15 + 20:
        result["解析提示"] = f"鉴权码长度声明 {auth_len} 字节后，剩余数据不足 IMEI 和软件版本号。"
        result[f"[{hex_slice(body)}]原始鉴权数据"] = safe_decode(body)
        return result
    auth_code, offset = read_bytes(body, offset, auth_len)
    imei, offset = read_bytes(body, offset, 15)
    software_version, offset = read_bytes(body, offset, 20)
    extra = body[offset:]
    result.update({
        f"[{auth_len:02X}]鉴权码长度": auth_len,
        f"[{hex_slice(auth_code)}]鉴权码内容": safe_decode(auth_code),
        f"[{hex_slice(imei)}]终端IMEI": safe_decode(imei),
        f"[{hex_slice(software_version)}]软件版本号": decode_trimmed_text(software_version),
    })
    if extra:
        result["扩展数据"] = hex_slice(extra)
    return result


def parse_bit_field(value, mapping):
    bits = {
        f"bit{bit}": bool((value >> bit) & 0x01)
        for bit in range(0, max(mapping.keys()) + 1 if mapping else 0)
    }
    labels = [name for bit, name in mapping.items() if (value >> bit) & 0x01]
    return bits, labels


def decode_trimmed_text(data):
    return safe_decode(data).rstrip("\x00").strip()


def parse_terminal_property_response(body, protocol="2013"):
    result = {"查询终端属性应答": hex_slice(body)}
    if len(body) < 48:
        result["解析提示"] = f"终端属性应答长度过短，当前为 {len(body)} 字节。"
        return result

    def try_layout(model_len, terminal_id_len, layout_name):
        offset = 0
        terminal_type, offset = read_u16(body, offset)
        manufacturer_id, offset = read_bytes(body, offset, 5)
        model, offset = read_bytes(body, offset, model_len)
        terminal_id, offset = read_bytes(body, offset, terminal_id_len)
        iccid, offset = read_bytes(body, offset, 10)
        hw_len, offset = read_u8(body, offset)
        hw_version, offset = read_bytes(body, offset, hw_len)
        fw_len, offset = read_u8(body, offset)
        fw_version, offset = read_bytes(body, offset, fw_len)
        gnss_attr, offset = read_u8(body, offset)
        comm_attr, offset = read_u8(body, offset)
        return {
            "layout_name": layout_name,
            "model_len": model_len,
            "terminal_id_len": terminal_id_len,
            "terminal_type": terminal_type,
            "manufacturer_id": manufacturer_id,
            "model": model,
            "terminal_id": terminal_id,
            "iccid": iccid,
            "hw_len": hw_len,
            "hw_version": hw_version,
            "fw_len": fw_len,
            "fw_version": fw_version,
            "gnss_attr": gnss_attr,
            "comm_attr": comm_attr,
            "extra": body[offset:],
            "offset": offset,
        }

    if protocol == "2019":
        layouts = [
            (30, 30, "JT/T 808-2019：型号30字节，终端ID30字节"),
        ]
    else:
        layouts = [
            (20, 15, "JT/T 808-2013：型号20字节，终端ID15字节"),
            (20, 7, "JT/T 808-2013 早期/窄终端ID兼容：型号20字节，终端ID7字节"),
        ]

    parsed = None
    parse_errors = []
    for model_len, terminal_id_len, layout_name in layouts:
        try:
            candidate = try_layout(model_len, terminal_id_len, layout_name)
            if not candidate["extra"]:
                parsed = candidate
                break
            if parsed is None or len(candidate["extra"]) < len(parsed["extra"]):
                parsed = candidate
        except ParseError as exc:
            parse_errors.append(f"{layout_name}: {exc}")

    if parsed is None:
        result["解析提示"] = "无法匹配终端属性应答字段长度布局。"
        result["尝试记录"] = parse_errors
        return result

    terminal_type = parsed["terminal_type"]
    manufacturer_id = parsed["manufacturer_id"]
    model = parsed["model"]
    terminal_id = parsed["terminal_id"]
    iccid = parsed["iccid"]
    hw_len = parsed["hw_len"]
    hw_version = parsed["hw_version"]
    fw_len = parsed["fw_len"]
    fw_version = parsed["fw_version"]
    gnss_attr = parsed["gnss_attr"]
    comm_attr = parsed["comm_attr"]
    extra = parsed["extra"]

    terminal_bits, terminal_labels = parse_bit_field(terminal_type, TERMINAL_TYPE_BIT_NAMES)
    gnss_bits, gnss_labels = parse_bit_field(gnss_attr, GNSS_MODULE_BIT_NAMES)
    comm_bits, comm_labels = parse_bit_field(comm_attr, COMM_MODULE_BIT_NAMES)

    result.update({
        "字段长度布局": parsed["layout_name"],
        f"[{body[0:2].hex().upper()}]终端类型": terminal_type,
        "终端类型对象": terminal_bits,
        "终端类型说明": terminal_labels,
        f"[{hex_slice(manufacturer_id)}]制造商ID": decode_trimmed_text(manufacturer_id),
        f"[{hex_slice(model)}]终端型号": decode_trimmed_text(model),
        f"[{hex_slice(terminal_id)}]终端ID": decode_trimmed_text(terminal_id),
        f"[{hex_slice(iccid)}]终端SIM卡ICCID": bytes_to_bcd(iccid),
        f"[{hw_len:02X}]终端硬件版本号长度": hw_len,
        f"[{hex_slice(hw_version)}]终端硬件版本号": decode_trimmed_text(hw_version),
        f"[{fw_len:02X}]终端固件版本号长度": fw_len,
        f"[{hex_slice(fw_version)}]终端固件版本号": decode_trimmed_text(fw_version),
        f"[{gnss_attr:02X}]GNSS模块属性": gnss_attr,
        "GNSS模块属性对象": gnss_bits,
        "GNSS模块属性说明": gnss_labels,
        f"[{comm_attr:02X}]通信模块属性": comm_attr,
        "通信模块属性对象": comm_bits,
        "通信模块属性说明": comm_labels,
    })
    if extra:
        result["扩展数据"] = hex_slice(extra)
    return result


def parse_upgrade_package(body):
    result = {"下发终端升级包": hex_slice(body)}
    if len(body) < 11:
        result["解析提示"] = f"标准升级包最短 11 字节，不含升级数据时也需要基本头部。当前为 {len(body)} 字节。"
        return result

    offset = 0
    upgrade_type, offset = read_u8(body, offset)
    manufacturer_id, offset = read_bytes(body, offset, 5)
    version_len, offset = read_u8(body, offset)
    version, offset = read_bytes(body, offset, version_len)
    package_len, offset = read_u32(body, offset)
    package_data = body[offset:offset + package_len]
    if len(package_data) != package_len:
        result["解析提示"] = f"升级数据包长度声明 {package_len} 字节，实际 {len(package_data)} 字节。"
        return result

    result.update({
        f"[{upgrade_type:02X}]升级类型": UPGRADE_TYPE_NAMES.get(upgrade_type, f"Unknown({upgrade_type})"),
        f"[{hex_slice(manufacturer_id)}]制造商ID": decode_trimmed_text(manufacturer_id),
        f"[{version_len:02X}]终端固件版本号长度": version_len,
        f"[{hex_slice(version)}]终端固件版本号": decode_trimmed_text(version),
        f"[{package_len:08X}]升级数据包长度": package_len,
        f"[{hex_slice(package_data)}]升级数据包": hex_slice(package_data),
    })
    return result


def parse_upgrade_result_notification(body):
    result = {"终端升级结果通知": hex_slice(body)}
    if len(body) < 2:
        result["解析提示"] = f"升级结果通知最短 2 字节，当前为 {len(body)} 字节。"
        return result
    upgrade_type, offset = read_u8(body, 0)
    upgrade_result, offset = read_u8(body, offset)
    result.update({
        f"[{upgrade_type:02X}]升级类型": UPGRADE_TYPE_NAMES.get(upgrade_type, f"Unknown({upgrade_type})"),
        f"[{upgrade_result:02X}]升级结果": UPGRADE_RESULT_NAMES.get(upgrade_result, f"Unknown({upgrade_result})"),
    })
    if offset < len(body):
        result["扩展数据"] = hex_slice(body[offset:])
    return result


def active_bit_names(value, names):
    active = {}
    for bit, name in names.items():
        active[f"bit{bit}"] = bool((value >> bit) & 0x01)
        if active[f"bit{bit}"]:
            active[f"bit{bit}说明"] = name
    return active


def format_number(value, digits=4):
    if isinstance(value, float):
        text = f"{value:.{digits}f}".rstrip("0").rstrip(".")
        return text or "0"
    return str(value)


def decode_e6_pid_value(pid, data):
    if data and all(item == 0xFF for item in data):
        return "无效", ""

    raw = int.from_bytes(data, "big") if data else 0
    rules = {
        1: ("仪表车速", "km/h", lambda x: x / 256),
        4: ("发动机实际扭矩百分比", "%", lambda x: x - 125),
        5: ("发动机转速", "rpm", lambda x: x * 0.125),
        6: ("冷却水温", "℃", lambda x: x - 40),
        8: ("ECU 总油耗", "L", lambda x: x * 0.5),
        9: ("油门踏板位置", "%", lambda x: x * 0.4),
        14: ("尿素箱液位", "%", lambda x: x * 0.4),
        20: ("发动机燃油消耗率", "L/h", lambda x: x * 0.05),
        21: ("发动机瞬时油耗率", "km/L", lambda x: x / 512),
        26: ("发动机进气压力", "kPa", lambda x: x * 2),
        31: ("发动机进气温度", "℃", lambda x: x - 40),
        39: ("仪表燃油液位", "%", lambda x: x * 0.4),
        41: ("仪表总里程", "km", lambda x: x * 0.125),
        42: ("积分里程", "km", lambda x: x / 1000),
        43: ("积分油耗", "L", lambda x: x * 0.0005),
        44: ("尿素实际喷射量", "ml/h", lambda x: x * 0.01),
        45: ("累计尿素消耗量", "g", lambda x: x),
    }
    bit_rules = {
        10: ("驻车制动开关", "", {0: "驻车制动开关在关闭状态", 1: "驻车制动开关在打开状态"}),
        12: ("离合开关状态", "", {0: "离合器踏板被松开", 1: "离合器踏板被踩下", 2: "错误", 3: "不可用"}),
    }
    if pid in rules:
        name, unit, converter = rules[pid]
        return format_number(converter(raw)), name + (f"({unit})" if unit else "")
    if pid in bit_rules:
        name, unit, states = bit_rules[pid]
        return states.get(raw, str(raw)), name + (f"({unit})" if unit else "")
    return hex_slice(data), f"未知PID({pid})"


def parse_e6_private_klv(value):
    result = {
        "E6原始值": hex_slice(value),
        "KLV列表": [],
        "摘要": "",
    }
    offset = 0
    summary_parts = []
    while offset < len(value):
        item_start = offset
        if offset + 3 > len(value):
            result["解析提示"] = f"E6 KLV 数据不足：偏移 {offset} 后不足 PID 和长度字段。"
            break
        pid_raw = value[offset:offset + 2]
        pid = int.from_bytes(pid_raw, "big")
        offset += 2
        data_len = value[offset]
        offset += 1
        if offset + data_len > len(value):
            result["解析提示"] = f"E6 PID {hex_slice(pid_raw)} 长度声明 {data_len} 字节，实际剩余 {len(value) - offset} 字节。"
            break
        data = value[offset:offset + data_len]
        offset += data_len
        decoded_value, label = decode_e6_pid_value(pid, data)
        pid_display = f"{pid:02X}"
        summary = f"{pid_display}-{decoded_value}-{label}" if label else f"{pid_display}-{decoded_value}"
        result["KLV列表"].append({
            "E6内部偏移": item_start,
            f"[{hex_slice(pid_raw)}]PID": pid,
            f"[{data_len:02X}]数据长度": data_len,
            f"[{hex_slice(data)}]数据": hex_slice(data),
            "解析值": decoded_value,
            "字段描述": label,
            "摘要": summary,
        })
        summary_parts.append(summary)
    result["摘要"] = "，".join(summary_parts)
    return result


def parse_location_extra(extra):
    result = {"附加信息原始": hex_slice(extra)}
    offset = 0
    items = []

    while offset < len(extra):
        item_start = offset
        item_id, offset = read_u8(extra, offset)
        item_len, offset = read_u8(extra, offset)
        require_length(extra, offset, item_len)
        value = extra[offset:offset + item_len]
        offset += item_len

        item = {
            f"[{item_id:02X}]附加信息ID": item_id,
            f"[{item_len:02X}]附加信息长度": item_len,
            "原始值": hex_slice(value),
        }
        if item_id == 0x01 and item_len == 4:
            item["里程(km)"] = int.from_bytes(value, "big") / 10
            item["摘要"] = f"{item_id:02X}-{format_number(item['里程(km)'])}-里程(km)"
        elif item_id == 0x02 and item_len == 2:
            item["油量(L)"] = int.from_bytes(value, "big") / 10
            item["摘要"] = f"{item_id:02X}-{format_number(item['油量(L)'])}-油量(L)"
        elif item_id == 0x03 and item_len == 2:
            item["行驶记录速度(km/h)"] = int.from_bytes(value, "big") / 10
            item["摘要"] = f"{item_id:02X}-{format_number(item['行驶记录速度(km/h)'])}-行驶记录速度(km/h)"
        elif item_id == 0x04 and item_len == 2:
            item["需要人工确认报警事件ID"] = int.from_bytes(value, "big")
            item["摘要"] = f"{item_id:02X}-{item['需要人工确认报警事件ID']}-需要人工确认报警事件ID"
        elif item_id == 0x25 and item_len == 4:
            item["扩展车辆信号状态"] = int.from_bytes(value, "big")
            item["摘要"] = f"{item_id:02X}-{item['扩展车辆信号状态']}-扩展车辆信号状态"
        elif item_id == 0x2A and item_len == 2:
            item["IO状态"] = int.from_bytes(value, "big")
            item["摘要"] = f"{item_id:02X}-{item['IO状态']}-IO状态"
        elif item_id == 0x30 and item_len == 1:
            item["无线通信网络信号强度"] = value[0]
            item["摘要"] = f"{item_id:02X}-{value[0]}-无线通信网络信号强度"
        elif item_id == 0x31 and item_len == 1:
            item["GNSS定位卫星数"] = value[0]
            item["摘要"] = f"{item_id:02X}-{value[0]}-GNSS定位卫星数"
        elif item_id == 0xE6:
            e6_obj = parse_e6_private_klv(value)
            item["E6私有协议对象"] = e6_obj
            item["摘要"] = e6_obj.get("摘要", "")
        else:
            item["解析提示"] = "未知或暂未细分的附加信息项"
            item["摘要"] = f"{item_id:02X}-{hex_slice(value)}"
        item["附加信息区偏移"] = item_start
        items.append(item)

    result["附加信息列表"] = items
    return result


def parse_location_report(body, name="位置信息汇报"):
    result = {name: hex_slice(body)}
    if len(body) < 28:
        result["解析提示"] = f"位置基础信息至少 28 字节，当前为 {len(body)} 字节。"
        return result

    offset = 0
    alarm, offset = read_u32(body, offset)
    status, offset = read_u32(body, offset)
    latitude_raw, offset = read_u32(body, offset)
    longitude_raw, offset = read_u32(body, offset)
    altitude, offset = read_u16(body, offset)
    speed_raw, offset = read_u16(body, offset)
    direction, offset = read_u16(body, offset)
    time_raw = body[offset:offset + 6]
    offset += 6
    extra = body[offset:]

    result.update({
        f"[{body[0:4].hex().upper()}]报警标志": alarm,
        "报警标志对象": active_bit_names(alarm, ALARM_BIT_NAMES),
        f"[{body[4:8].hex().upper()}]状态": status,
        "状态对象": active_bit_names(status, STATUS_BIT_NAMES),
        f"[{body[8:12].hex().upper()}]纬度": latitude_raw / 1_000_000,
        f"[{body[12:16].hex().upper()}]经度": longitude_raw / 1_000_000,
        f"[{body[16:18].hex().upper()}]高程(m)": altitude,
        f"[{body[18:20].hex().upper()}]速度(km/h)": speed_raw / 10,
        f"[{body[20:22].hex().upper()}]方向": direction,
        f"[{hex_slice(time_raw)}]时间": bytes_to_bcd(time_raw),
    })
    if extra:
        result["附加信息对象"] = parse_location_extra(extra)
    return result


def parse_register_response(body):
    result = {"终端注册应答": hex_slice(body)}
    if len(body) < 3:
        result["解析提示"] = f"注册应答最短 3 字节，当前为 {len(body)} 字节。"
        return result
    response_serial = int.from_bytes(body[0:2], "big")
    register_result = body[2]
    result.update({
        f"[{body[0:2].hex().upper()}]应答流水号": response_serial,
        f"[{body[2]:02X}]结果": REGISTER_RESULT_NAMES.get(register_result, f"Unknown({register_result})"),
    })
    if len(body) > 3:
        auth = body[3:]
        result[f"[{hex_slice(auth)}]鉴权码"] = safe_decode(auth)
    return result


def parse_query_param_response(body):
    result = {"查询终端参数应答": hex_slice(body)}
    if len(body) < 3:
        result["解析提示"] = f"参数应答最短 3 字节，当前为 {len(body)} 字节。"
        return result

    offset = 0
    response_serial, offset = read_u16(body, offset)
    count, offset = read_u8(body, offset)
    params = []
    for _ in range(count):
        if offset >= len(body):
            break
        param_id, offset = read_u32(body, offset)
        param_len, offset = read_u8(body, offset)
        require_length(body, offset, param_len)
        value = body[offset:offset + param_len]
        offset += param_len
        params.append({
            f"[{param_id:08X}]参数ID": param_id,
            f"[{param_len:02X}]参数长度": param_len,
            "参数值HEX": hex_slice(value),
            "参数值文本": safe_decode(value),
        })

    result.update({
        f"[{body[0:2].hex().upper()}]应答流水号": response_serial,
        f"[{body[2]:02X}]应答参数个数": count,
        "参数列表": params,
    })
    if len(params) != count:
        result["解析提示"] = f"声明 {count} 个参数，实际解析 {len(params)} 个。"
    return result


def parse_set_params(body):
    result = {"设置终端参数": hex_slice(body)}
    if not body:
        result["解析提示"] = "参数个数为空。"
        return result

    offset = 0
    count, offset = read_u8(body, offset)
    params = []
    for _ in range(count):
        if offset >= len(body):
            break
        param_id, offset = read_u32(body, offset)
        param_len, offset = read_u8(body, offset)
        require_length(body, offset, param_len)
        value = body[offset:offset + param_len]
        offset += param_len
        params.append({
            f"[{param_id:08X}]参数ID": param_id,
            f"[{param_len:02X}]参数长度": param_len,
            "参数值HEX": hex_slice(value),
            "参数值文本": safe_decode(value),
        })

    result.update({
        f"[{body[0]:02X}]参数总数": count,
        "参数列表": params,
    })
    if len(params) != count:
        result["解析提示"] = f"声明 {count} 个参数，实际解析 {len(params)} 个。"
    return result


def parse_retransmission_request(body, name):
    result = {name: hex_slice(body)}
    if len(body) < 3:
        result["解析提示"] = f"补传分包请求最短 3 字节，当前为 {len(body)} 字节。"
        return result
    offset = 0
    original_serial, offset = read_u16(body, offset)
    count, offset = read_u8(body, offset)
    packet_ids = []
    for _ in range(count):
        if offset + 2 > len(body):
            break
        packet_id, offset = read_u16(body, offset)
        packet_ids.append(packet_id)
    result.update({
        f"[{body[0:2].hex().upper()}]原始消息流水号": original_serial,
        f"[{body[2]:02X}]重传包总数": count,
        "重传包ID列表": packet_ids,
    })
    if len(packet_ids) != count:
        result["解析提示"] = f"声明 {count} 个重传包 ID，实际解析 {len(packet_ids)} 个。"
    return result


def parse_query_specific_params(body):
    result = {"查询指定终端参数": hex_slice(body)}
    if not body:
        result["解析提示"] = "缺少参数总数。"
        return result
    offset = 0
    count, offset = read_u8(body, offset)
    param_ids = []
    for _ in range(count):
        if offset + 4 > len(body):
            break
        param_id, offset = read_u32(body, offset)
        param_ids.append(f"0x{param_id:08X}")
    result.update({
        f"[{body[0]:02X}]参数总数": count,
        "参数ID列表": param_ids,
    })
    if len(param_ids) != count:
        result["解析提示"] = f"声明 {count} 个参数 ID，实际解析 {len(param_ids)} 个。"
    return result


def parse_terminal_control(body):
    result = {"终端控制": hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少命令字。"
        return result
    command_word = body[0]
    command_param = body[1:]
    result.update({
        f"[{command_word:02X}]命令字": command_word,
        f"[{hex_slice(command_param)}]命令参数": decode_trimmed_text(command_param),
    })
    return result


def parse_server_time_response(body):
    result = {"查询服务器时间应答": hex_slice(body)}
    if len(body) != 6:
        result["解析提示"] = f"服务器时间应答标准长度为 6 字节 BCD，当前为 {len(body)} 字节。"
        return result
    result[f"[{hex_slice(body)}]UTC时间"] = "20" + bytes_to_bcd(body)
    return result


def parse_text_info(body):
    result = {"文本信息下发": hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少标志字段。"
        return result
    flag = body[0]
    text = body[1:]
    result.update({
        f"[{flag:02X}]标志": flag,
        "标志对象": {
            "bit0紧急": bool(flag & 0x01),
            "bit2终端显示器显示": bool(flag & 0x04),
            "bit3终端TTS播读": bool(flag & 0x08),
            "bit4广告屏显示": bool(flag & 0x10),
        },
        f"[{hex_slice(text)}]文本信息": safe_decode(text),
    })
    return result


def parse_event_setting(body):
    result = {"事件设置": hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少设置类型。"
        return result
    offset = 0
    setting_type, offset = read_u8(body, offset)
    result[f"[{setting_type:02X}]设置类型"] = setting_type
    if offset >= len(body):
        return result

    count, offset = read_u8(body, offset)
    events = []
    for _ in range(count):
        if offset + 2 > len(body):
            break
        event_id, offset = read_u8(body, offset)
        event_len, offset = read_u8(body, offset)
        if offset + event_len > len(body):
            break
        event_content, offset = read_bytes(body, offset, event_len)
        events.append({
            f"[{event_id:02X}]事件ID": event_id,
            f"[{event_len:02X}]事件内容长度": event_len,
            f"[{hex_slice(event_content)}]事件内容": safe_decode(event_content),
        })
    result[f"[{count:02X}]事件总数"] = count
    result["事件列表"] = events
    if len(events) != count:
        result["解析提示"] = f"声明 {count} 个事件，实际解析 {len(events)} 个。"
    return result


def parse_info_demand_cancel(body):
    result = {"信息点播/取消": hex_slice(body)}
    if len(body) < 2:
        result["解析提示"] = f"标准长度为 2 字节，当前为 {len(body)} 字节。"
        return result
    result.update({
        f"[{body[0]:02X}]信息类型": body[0],
        f"[{body[1]:02X}]点播/取消标志": "取消" if body[1] == 0 else "点播",
    })
    return result


def parse_info_service(body):
    result = {"信息服务": hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少信息类型。"
        return result
    result.update({
        f"[{body[0]:02X}]信息类型": body[0],
        f"[{hex_slice(body[1:])}]信息内容": safe_decode(body[1:]),
    })
    return result


def parse_phone_callback(body):
    result = {"电话回拨": hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少标志字段。"
        return result
    result.update({
        f"[{body[0]:02X}]标志": "普通通话" if body[0] == 0 else "监听",
        f"[{hex_slice(body[1:])}]电话号码": safe_decode(body[1:]),
    })
    return result


def parse_phonebook_setting(body):
    result = {"设置电话本": hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少设置类型。"
        return result
    offset = 0
    setting_type, offset = read_u8(body, offset)
    result[f"[{setting_type:02X}]设置类型"] = PHONEBOOK_SETTING_TYPE_NAMES.get(setting_type, f"Unknown({setting_type})")
    if offset >= len(body):
        return result
    count, offset = read_u8(body, offset)
    contacts = []
    for _ in range(count):
        if offset + 3 > len(body):
            break
        flag, offset = read_u8(body, offset)
        phone_len, offset = read_u8(body, offset)
        if offset + phone_len + 1 > len(body):
            break
        phone, offset = read_bytes(body, offset, phone_len)
        contact_len, offset = read_u8(body, offset)
        if offset + contact_len > len(body):
            break
        contact, offset = read_bytes(body, offset, contact_len)
        contacts.append({
            f"[{flag:02X}]标志": flag,
            f"[{phone_len:02X}]电话号码长度": phone_len,
            f"[{hex_slice(phone)}]电话号码": safe_decode(phone),
            f"[{contact_len:02X}]联系人长度": contact_len,
            f"[{hex_slice(contact)}]联系人": safe_decode(contact),
        })
    result[f"[{count:02X}]联系人总数"] = count
    result["联系人列表"] = contacts
    if len(contacts) != count:
        result["解析提示"] = f"声明 {count} 个联系人，实际解析 {len(contacts)} 个。"
    return result


def parse_vehicle_control(body):
    result = {"车辆控制": hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少控制标志。"
        return result
    result.update({
        f"[{body[0]:02X}]控制标志": body[0],
        "控制标志对象": {
            "bit0车门": "车门加锁" if body[0] & 0x01 else "车门解锁",
        },
    })
    return result


def parse_vehicle_control_response(body):
    result = {"车辆控制应答": hex_slice(body)}
    if len(body) < 2:
        result["解析提示"] = "缺少应答流水号。"
        return result
    result[f"[{body[0:2].hex().upper()}]应答流水号"] = int.from_bytes(body[0:2], "big")
    result["位置数据对象"] = parse_location_report(body[2:], "位置数据") if len(body) > 2 else {}
    return result


def parse_alarm_manual_confirm(body):
    result = {"人工确认报警消息": hex_slice(body)}
    if len(body) < 6:
        result["解析提示"] = f"标准长度为 6 字节，当前为 {len(body)} 字节。"
        return result
    result.update({
        f"[{body[0:2].hex().upper()}]报警消息流水号": int.from_bytes(body[0:2], "big"),
        f"[{body[2:6].hex().upper()}]人工确认报警类型": int.from_bytes(body[2:6], "big"),
    })
    return result


def parse_temp_location_tracking(body):
    result = {"临时位置跟踪控制": hex_slice(body)}
    if len(body) < 6:
        result["解析提示"] = f"标准长度至少 6 字节，当前为 {len(body)} 字节。"
        return result
    interval = int.from_bytes(body[0:2], "big")
    validity = int.from_bytes(body[2:6], "big")
    result.update({
        f"[{body[0:2].hex().upper()}]时间间隔(s)": interval,
        f"[{body[2:6].hex().upper()}]位置跟踪有效期(s)": validity,
    })
    return result


def parse_batch_location(body):
    result = {"定位数据批量上传": hex_slice(body)}
    if len(body) < 3:
        result["解析提示"] = f"批量位置最短 3 字节，当前为 {len(body)} 字节。"
        return result

    offset = 0
    count, offset = read_u16(body, offset)
    data_type, offset = read_u8(body, offset)
    locations = []
    for index in range(count):
        if offset >= len(body):
            break
        location_len, offset = read_u16(body, offset)
        require_length(body, offset, location_len)
        location_body = body[offset:offset + location_len]
        offset += location_len
        locations.append({
            "序号": index + 1,
            f"[{location_len:04X}]位置数据体长度": location_len,
            "位置数据体对象": parse_location_report(location_body, "位置数据体"),
        })

    result.update({
        f"[{body[0:2].hex().upper()}]数据项个数": count,
        f"[{body[2]:02X}]位置数据类型": "正常位置批量汇报" if data_type == 0 else f"盲区补报({data_type})",
        "位置数据项列表": locations,
    })
    if len(locations) != count:
        result["解析提示"] = f"声明 {count} 条位置数据，实际解析 {len(locations)} 条。"
    return result


def parse_can_bus_upload(body):
    result = {"CAN总线数据上传": hex_slice(body)}
    if len(body) < 7:
        result["解析提示"] = f"CAN 数据上传最短 7 字节，当前为 {len(body)} 字节。"
        return result
    offset = 0
    count, offset = read_u16(body, offset)
    time_raw, offset = read_bytes(body, offset, 5)
    items = []
    for index in range(count):
        if offset + 12 > len(body):
            break
        can_id, offset = read_u32(body, offset)
        can_data, offset = read_bytes(body, offset, 8)
        items.append({
            "序号": index + 1,
            f"[{can_id:08X}]CAN ID": can_id,
            f"[{hex_slice(can_data)}]CAN数据": hex_slice(can_data),
        })
    result.update({
        f"[{body[0:2].hex().upper()}]数据项个数": count,
        f"[{hex_slice(time_raw)}]CAN数据接收时间": bytes_to_bcd(time_raw),
        "CAN数据列表": items,
    })
    if len(items) != count:
        result["解析提示"] = f"声明 {count} 条 CAN 数据，实际解析 {len(items)} 条。"
    return result


def parse_multimedia_event(body):
    result = {"多媒体事件信息上传": hex_slice(body)}
    if len(body) < 8:
        result["解析提示"] = f"多媒体事件信息标准长度为 8 字节，当前为 {len(body)} 字节。"
        return result
    result.update({
        f"[{body[0:4].hex().upper()}]多媒体数据ID": int.from_bytes(body[0:4], "big"),
        f"[{body[4]:02X}]多媒体类型": body[4],
        f"[{body[5]:02X}]多媒体格式编码": body[5],
        f"[{body[6]:02X}]事件项编码": body[6],
        f"[{body[7]:02X}]通道ID": body[7],
    })
    if len(body) > 8:
        result["扩展数据"] = hex_slice(body[8:])
    return result


def parse_multimedia_upload(body):
    result = {"多媒体数据上传": hex_slice(body)}
    if len(body) < 36:
        result["解析提示"] = f"多媒体数据上传至少包含 8 字节多媒体头 + 28 字节位置基础信息，当前为 {len(body)} 字节。"
        return result
    result.update(parse_multimedia_event(body[:8]))
    result["位置数据对象"] = parse_location_report(body[8:36], "位置数据")
    result[f"[{hex_slice(body[36:])}]多媒体数据包"] = hex_slice(body[36:])
    return result


def parse_multimedia_upload_response(body):
    result = {"多媒体数据上传应答": hex_slice(body)}
    if len(body) < 5:
        result["解析提示"] = f"多媒体数据上传应答最短 5 字节，当前为 {len(body)} 字节。"
        return result
    offset = 0
    multimedia_id, offset = read_u32(body, offset)
    count, offset = read_u8(body, offset)
    packet_ids = []
    for _ in range(count):
        if offset + 2 > len(body):
            break
        packet_id, offset = read_u16(body, offset)
        packet_ids.append(packet_id)
    result.update({
        f"[{multimedia_id:08X}]多媒体数据ID": multimedia_id,
        f"[{count:02X}]重传包总数": count,
        "重传包ID列表": packet_ids,
    })
    if len(packet_ids) != count:
        result["解析提示"] = f"声明 {count} 个重传包 ID，实际解析 {len(packet_ids)} 个。"
    return result


def parse_camera_shoot_response(body):
    result = {"摄像头立即拍摄命令应答": hex_slice(body)}
    if len(body) < 3:
        result["解析提示"] = f"拍摄命令应答最短 3 字节，当前为 {len(body)} 字节。"
        return result
    offset = 0
    response_serial, offset = read_u16(body, offset)
    shoot_result, offset = read_u8(body, offset)
    multimedia_ids = []
    while offset + 4 <= len(body):
        multimedia_id, offset = read_u32(body, offset)
        multimedia_ids.append(multimedia_id)
    result.update({
        f"[{body[0:2].hex().upper()}]应答流水号": response_serial,
        f"[{shoot_result:02X}]结果": shoot_result,
        "多媒体ID列表": multimedia_ids,
    })
    return result


def parse_delete_id_list(body, name):
    result = {name: hex_slice(body)}
    if len(body) < 1:
        result["解析提示"] = "缺少区域/路线数。"
        return result
    offset = 0
    count, offset = read_u8(body, offset)
    ids = []
    for _ in range(count):
        if offset + 4 > len(body):
            break
        item_id, offset = read_u32(body, offset)
        ids.append(item_id)
    result.update({
        f"[{count:02X}]区域/路线总数": count,
        "ID列表": [f"0x{item_id:08X}" for item_id in ids],
    })
    if len(ids) != count:
        result["解析提示"] = f"声明 {count} 个 ID，实际解析 {len(ids)} 个。"
    return result


def parse_raw_with_note(body, name, note):
    return {
        name: hex_slice(body),
        "解析提示": note,
    }


def parse_empty_body(body, name):
    if body:
        return {name: hex_slice(body), "解析提示": f"该标准消息体应为空，当前存在 {len(body)} 字节数据。"}
    return {name: "", "说明": "空消息体"}


def parse_body(msg_id, body, protocol="2013"):
    if msg_id == 0x0001:
        return parse_common_response(body, "终端通用应答")
    if msg_id == 0x8001:
        return parse_common_response(body, "平台通用应答")
    if msg_id in (0x0002, 0x0003, 0x0004, 0x8104, 0x8107, 0x8201, 0x8204, 0x8702):
        return parse_empty_body(body, MESSAGE_ID_NAMES.get(msg_id, f"Unknown(0x{msg_id:04X})"))
    if msg_id == 0x0005:
        return parse_retransmission_request(body, "终端补传分包请求")
    if msg_id == 0x0100:
        return parse_terminal_register(body, protocol)
    if msg_id == 0x0102:
        return parse_authentication(body, protocol)
    if msg_id == 0x0104:
        return parse_query_param_response(body)
    if msg_id == 0x0107:
        return parse_terminal_property_response(body, protocol)
    if msg_id == 0x0108:
        return parse_upgrade_result_notification(body)
    if msg_id == 0x0200:
        return parse_location_report(body)
    if msg_id == 0x0201:
        result = {"位置信息查询应答": hex_slice(body)}
        if len(body) >= 2:
            result[f"[{body[0:2].hex().upper()}]应答流水号"] = int.from_bytes(body[0:2], "big")
            result["位置数据对象"] = parse_location_report(body[2:], "位置数据")
        else:
            result["解析提示"] = "缺少应答流水号。"
        return result
    if msg_id == 0x0301:
        if len(body) == 1:
            return {"事件报告": hex_slice(body), f"[{body[0]:02X}]事件ID": body[0]}
        return {"事件报告": hex_slice(body), "解析提示": f"标准事件报告长度为 1 字节，当前为 {len(body)} 字节。"}
    if msg_id == 0x0302:
        result = {"提问应答": hex_slice(body)}
        if len(body) >= 3:
            result[f"[{body[0:2].hex().upper()}]应答流水号"] = int.from_bytes(body[0:2], "big")
            result[f"[{body[2]:02X}]答案ID"] = body[2]
        else:
            result["解析提示"] = f"提问应答最短 3 字节，当前为 {len(body)} 字节。"
        return result
    if msg_id == 0x0303:
        return parse_info_demand_cancel(body)
    if msg_id == 0x0500:
        return parse_vehicle_control_response(body)
    if msg_id == 0x0700:
        return parse_raw_with_note(body, "行驶记录仪数据上传", "已识别标准消息 ID，消息体按行驶记录仪透传数据保留原始 HEX。")
    if msg_id == 0x0702:
        return parse_raw_with_note(body, "驾驶员身份信息采集上报", "已识别标准消息 ID，驾驶员身份信息字段受状态和证件数据影响，当前保留原始 HEX。")
    if msg_id == 0x0704:
        return parse_batch_location(body)
    if msg_id == 0x0705:
        return parse_can_bus_upload(body)
    if msg_id == 0x0800:
        return parse_multimedia_event(body)
    if msg_id == 0x0801:
        return parse_multimedia_upload(body)
    if msg_id == 0x8003:
        return parse_retransmission_request(body, "服务器补传分包请求")
    if msg_id == 0x8004:
        return parse_server_time_response(body)
    if msg_id == 0x8100:
        return parse_register_response(body)
    if msg_id == 0x8103:
        return parse_set_params(body)
    if msg_id == 0x8105:
        return parse_terminal_control(body)
    if msg_id == 0x8106:
        return parse_query_specific_params(body)
    if msg_id == 0x8108:
        return parse_upgrade_package(body)
    if msg_id == 0x8202:
        return parse_temp_location_tracking(body)
    if msg_id == 0x8203:
        return parse_alarm_manual_confirm(body)
    if msg_id == 0x8300:
        return parse_text_info(body)
    if msg_id == 0x8301:
        return parse_event_setting(body)
    if msg_id == 0x8302:
        return parse_raw_with_note(body, "提问下发", "已识别标准消息 ID，可变候选答案列表暂按原始 HEX 保留。")
    if msg_id == 0x8303:
        return parse_raw_with_note(body, "信息点播菜单设置", "已识别标准消息 ID，可变菜单项暂按原始 HEX 保留。")
    if msg_id == 0x8304:
        return parse_info_service(body)
    if msg_id == 0x8400:
        return parse_phone_callback(body)
    if msg_id == 0x8401:
        return parse_phonebook_setting(body)
    if msg_id == 0x8500:
        return parse_vehicle_control(body)
    if msg_id in (0x8601, 0x8603, 0x8605, 0x8607):
        return parse_delete_id_list(body, MESSAGE_ID_NAMES.get(msg_id, f"Unknown(0x{msg_id:04X})"))
    if msg_id in (0x8600, 0x8602, 0x8604, 0x8606):
        return parse_raw_with_note(body, MESSAGE_ID_NAMES.get(msg_id, f"Unknown(0x{msg_id:04X})"), "已识别标准区域/路线设置消息，字段为可变结构，当前保留原始 HEX。")
    if msg_id in (0x8700, 0x8701):
        return parse_raw_with_note(body, MESSAGE_ID_NAMES.get(msg_id, f"Unknown(0x{msg_id:04X})"), "已识别标准行驶记录仪消息，当前保留原始 HEX。")
    if msg_id == 0x8800:
        return parse_multimedia_upload_response(body)
    if msg_id == 0x8801:
        return parse_raw_with_note(body, "摄像头立即拍摄命令", "已识别标准消息 ID，当前保留拍摄参数原始 HEX。")
    if msg_id == 0x0805:
        return parse_camera_shoot_response(body)

    return {
        MESSAGE_ID_NAMES.get(msg_id, f"Unknown(0x{msg_id:04X})"): hex_slice(body),
        "解析提示": "该消息 ID 暂未做消息体细分，已保留原始 HEX。",
    }


def normalize_protocol(protocol):
    protocol_text = str(protocol or "2013")
    return "2019" if "2019" in protocol_text else "2013"


def parse_one_frame(frame, protocol="2013"):
    protocol = normalize_protocol(protocol)
    if len(frame) < 2:
        raise ParseError("帧长度过短。")
    if frame[0] != 0x7E or frame[-1] != 0x7E:
        raise ParseError("标准 808 帧必须以 7E 开始并以 7E 结束。")

    escaped_payload = frame[1:-1]
    payload = unescape_808(escaped_payload)
    if len(payload) < 13:
        raise ParseError("帧内容过短，无法包含标准消息头和校验码。")

    packet_without_checksum = payload[:-1]
    received_checksum = payload[-1]
    calculated_checksum = xor_checksum(packet_without_checksum)
    msg_id, body, header = parse_header(packet_without_checksum, protocol)

    result = {"[7E]开始": 126}
    result.update(header)
    result["数据体对象"] = parse_body(msg_id, body, protocol)
    result[f"[{received_checksum:02X}]校验码"] = received_checksum
    result["校验结果对象"] = {
        "接收校验码": received_checksum,
        "计算校验码": calculated_checksum,
        "是否通过": received_checksum == calculated_checksum,
    }
    result["[7E]结束"] = 126
    return result


def split_frames(data):
    if not data:
        return []
    if data[0] == 0x7E and data[-1] == 0x7E and data.count(0x7E) == 2:
        return [data]

    frames = []
    start = None
    for index, byte in enumerate(data):
        if byte == 0x7E:
            if start is None:
                start = index
            else:
                if index > start:
                    frames.append(data[start:index + 1])
                start = index
    return frames


def parse_808_hex(hex_text, protocol="2013"):
    protocol = normalize_protocol(protocol)
    cleaned, data = ensure_hex_bytes(hex_text)
    frames = split_frames(data)
    if not frames:
        raise ParseError("未找到完整的 7E...7E 帧。")

    parsed_frames = [parse_one_frame(frame, protocol) for frame in frames]
    if len(parsed_frames) == 1 and len(frames[0]) == len(data):
        return parsed_frames[0]
    return {
        "清洗后HEX": cleaned,
        "帧数量": len(parsed_frames),
        "解析结果列表": parsed_frames,
    }


def extract_zipsend_hex(line):
    match = re.search(r"AT\+ZIPSEND\s*=\s*[^,]*,\s*([0-9A-Fa-f\s]+)", line, re.IGNORECASE)
    if not match:
        return None
    candidate = clean_hex_text(match.group(1))
    return candidate or None


def has_parse_warning(value):
    warning_keywords = ("解析提示", "解析错误", "解析失败", "数据长度不足", "无法匹配", "非法")
    if isinstance(value, dict):
        for key, item in value.items():
            if any(keyword in str(key) for keyword in warning_keywords):
                return True
            if has_parse_warning(item):
                return True
    elif isinstance(value, list):
        return any(has_parse_warning(item) for item in value)
    elif isinstance(value, str):
        return any(keyword in value for keyword in warning_keywords)
    return False


def find_top_value_by_suffix(data, suffix):
    for key, value in data.items():
        if isinstance(key, str) and key.endswith(suffix):
            return key, value
    return "", ""


def find_body_value_by_name(body_obj, name):
    if not isinstance(body_obj, dict):
        return ""
    for key, value in body_obj.items():
        if isinstance(key, str) and key.endswith(name):
            return value
    return ""


def build_location_extra_summary(body_obj):
    extra_obj = body_obj.get("附加信息对象", {}) if isinstance(body_obj, dict) else {}
    items = extra_obj.get("附加信息列表", []) if isinstance(extra_obj, dict) else []
    normal_parts = []
    e6_parts = []
    for item in items:
        if not isinstance(item, dict):
            continue
        item_id = None
        for key, value in item.items():
            if isinstance(key, str) and key.endswith("附加信息ID"):
                item_id = value
                break
        summary = item.get("摘要", "")
        if item_id == 0xE6:
            e6_obj = item.get("E6私有协议对象", {})
            e6_parts.append(e6_obj.get("摘要", summary))
        elif summary:
            normal_parts.append(summary)
    return "，".join(part for part in normal_parts if part), "，".join(part for part in e6_parts if part)


def build_sheet1_row(index, raw_hex, protocol, parsed=None, error_text=""):
    if parsed is None:
        return [
            index,
            "",
            "",
            "",
            "JT/T 808-2019" if normalize_protocol(protocol) == "2019" else "JT/T 808-2013",
            "",
            "",
            "",
            "",
            f"解析失败：{error_text}",
            "",
            "",
            raw_hex,
        ]

    msg_id_key, msg_id = find_top_value_by_suffix(parsed, "消息Id")
    phone_key, phone = find_top_value_by_suffix(parsed, "终端手机号")
    serial_key, serial = find_top_value_by_suffix(parsed, "消息流水号")
    checksum_key, checksum = find_top_value_by_suffix(parsed, "校验码")
    prop_obj = parsed.get("消息体属性对象", {})
    body_obj = parsed.get("数据体对象", {})

    msg_id_text = f"0x{msg_id:04X}" if isinstance(msg_id, int) else ""
    checksum_text = f"0x{checksum:02X}" if isinstance(checksum, int) else str(checksum or "")
    prop_text = json.dumps(prop_obj, ensure_ascii=False, separators=(",", ":"))
    body_text = json.dumps(body_obj, ensure_ascii=False, separators=(",", ":"))

    return [
        index,
        "7E" if parsed.get("[7E]开始") == 126 else "",
        msg_id_text,
        parsed.get("消息名称", ""),
        parsed.get("解析标准", "JT/T 808-2019" if normalize_protocol(protocol) == "2019" else "JT/T 808-2013"),
        prop_text,
        prop_obj.get("[bit10~bit12]数据加密", ""),
        phone,
        serial,
        body_text,
        checksum_text,
        "7E" if parsed.get("[7E]结束") == 126 else "",
        raw_hex,
    ]


def build_location_row(parsed):
    body_obj = parsed.get("数据体对象", {})
    extra_summary, e6_summary = build_location_extra_summary(body_obj)
    return [
        find_body_value_by_name(body_obj, "纬度"),
        find_body_value_by_name(body_obj, "经度"),
        find_body_value_by_name(body_obj, "高程(m)"),
        find_body_value_by_name(body_obj, "速度(km/h)"),
        find_body_value_by_name(body_obj, "方向"),
        find_body_value_by_name(body_obj, "时间"),
        extra_summary,
        e6_summary,
    ]


def autofit_worksheet(ws, max_width=80):
    for column_cells in ws.columns:
        column_letter = column_cells[0].column_letter
        width = 10
        for cell in column_cells:
            if cell.value is not None:
                width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[column_letter].width = width


def batch_parse_zipsend_file(input_path, output_path, protocol="2013"):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ParseError("缺少 openpyxl 库，无法生成 Excel。请先安装 openpyxl。") from exc

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "sheet1"
    ws2 = wb.create_sheet("location")

    sheet1_headers = [
        "序号", "消息头", "消息ID", "消息名称", "解析标准", "消息体属性", "数据加密",
        "手机号", "消息流水号", "消息体对象", "校验码", "消息尾", "完整HEX数据",
    ]
    location_headers = ["纬度", "经度", "高度", "速度", "方向", "时间", "附加信息列表", "E6数据"]
    ws1.append(sheet1_headers)
    ws2.append(location_headers)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    error_fill = PatternFill("solid", fgColor="FFC7CE")
    header_font = Font(bold=True)
    wrap_alignment = Alignment(vertical="top", wrap_text=True)

    for ws in (ws1, ws2):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"

    parsed_count = 0
    error_count = 0
    location_count = 0

    with open(input_path, "r", encoding="utf-8", errors="ignore") as file:
        for line in file:
            raw_hex = extract_zipsend_hex(line)
            if not raw_hex:
                continue

            parsed_count += 1
            row_has_error = False
            try:
                parsed = parse_808_hex(raw_hex, protocol)
                if "解析结果列表" in parsed:
                    parsed = parsed["解析结果列表"][0]
                row_has_error = has_parse_warning(parsed) or not parsed.get("校验结果对象", {}).get("是否通过", False)
                row = build_sheet1_row(parsed_count, raw_hex, protocol, parsed=parsed)
                msg_id = find_top_value_by_suffix(parsed, "消息Id")[1]
                if msg_id == 0x0200:
                    ws2.append(build_location_row(parsed))
                    location_count += 1
            except Exception as exc:
                row_has_error = True
                row = build_sheet1_row(parsed_count, raw_hex, protocol, error_text=str(exc))

            ws1.append(row)
            current_row = ws1.max_row
            for cell in ws1[current_row]:
                cell.alignment = wrap_alignment
                if row_has_error:
                    cell.fill = error_fill
            if row_has_error:
                error_count += 1

    for ws in (ws1, ws2):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap_alignment
        autofit_worksheet(ws)

    wb.save(output_path)
    return {
        "input_path": input_path,
        "output_path": output_path,
        "parsed_count": parsed_count,
        "error_count": error_count,
        "location_count": location_count,
    }


class JT808ParserApp:
    def __init__(self, root):
        self.root = root
        self.root.title("GT117-YK-808 标准808协议解析工具-Mangetic-260804")
        self.root.geometry("1060x720")
        self.root.minsize(860, 560)
        self._build_ui()
        self.load_sample()

    def _build_ui(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        main = ttk.Frame(self.root, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        toolbar = ttk.Frame(main)
        toolbar.pack(fill=tk.X, pady=(8, 8))

        ttk.Label(toolbar, text="协议标准").pack(side=tk.LEFT, padx=(0, 4))
        self.protocol_var = tk.StringVar(value="JT/T 808-2013")
        self.protocol_combo = ttk.Combobox(
            toolbar,
            textvariable=self.protocol_var,
            values=("JT/T 808-2013", "JT/T 808-2019"),
            state="readonly",
            width=16,
        )
        self.protocol_combo.pack(side=tk.LEFT, padx=(0, 8))
        self.protocol_combo.bind("<<ComboboxSelected>>", lambda _event: self.parse_input())

        ttk.Button(toolbar, text="解析", command=self.parse_input).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(toolbar, text="示例", command=self.load_sample).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(toolbar, text="清空输入", command=self.clear_input).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(toolbar, text="复制结果", command=self.copy_output).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(toolbar, text="批量解析文件", command=self.batch_parse_file).pack(side=tk.LEFT, padx=(0, 8))

        self.status_var = tk.StringVar(value="请粘贴 808 协议 HEX 数据，工具会自动去掉回车、换行、空格后解析。")
        ttk.Label(toolbar, textvariable=self.status_var).pack(side=tk.LEFT, padx=(8, 0))

        panes = ttk.Panedwindow(main, orient=tk.HORIZONTAL)
        panes.pack(fill=tk.BOTH, expand=True)

        input_frame = ttk.LabelFrame(panes, text="输入 HEX")
        output_frame = ttk.LabelFrame(panes, text="解析结果 JSON")
        panes.add(input_frame, weight=1)
        panes.add(output_frame, weight=1)

        self.input_text = scrolledtext.ScrolledText(input_frame, wrap=tk.WORD, font=("Consolas", 11), undo=True)
        self.input_text.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        self.output_text = scrolledtext.ScrolledText(output_frame, wrap=tk.WORD, font=("Consolas", 10), undo=True)
        self.output_text.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.output_text.tag_configure("error", foreground="#B00020", font=("Consolas", 10, "bold"))
        self.output_text.tag_configure("hint", foreground="#7A4A00", font=("Consolas", 10))

    def show_output(self, content):
        self.output_text.delete("1.0", tk.END)
        error_keywords = ("解析失败", "未知错误", "解析错误", "解析提示", "数据长度不足", "无法匹配", "非法", "不足")
        warning_keywords = ("尝试记录", "建议")
        for line in content.splitlines(True):
            if any(keyword in line for keyword in error_keywords):
                self.output_text.insert(tk.END, line, "error")
            elif any(keyword in line for keyword in warning_keywords):
                self.output_text.insert(tk.END, line, "hint")
            else:
                self.output_text.insert(tk.END, line)

    def load_sample(self):
        self.input_text.delete("1.0", tk.END)
        self.input_text.insert(tk.END, SAMPLE_HEX)
        self.parse_input()

    def clear_input(self):
        self.input_text.delete("1.0", tk.END)
        self.output_text.delete("1.0", tk.END)
        self.status_var.set("已清空。")

    def parse_input(self):
        source = self.input_text.get("1.0", tk.END)
        try:
            result = parse_808_hex(source, self.protocol_var.get())
            formatted = json.dumps(result, ensure_ascii=False, indent=2)
            self.show_output(formatted)
            if "解析提示" in formatted or "解析错误" in formatted or "数据长度不足" in formatted:
                self.status_var.set(f"解析完成，但存在红色提示；当前标准：{self.protocol_var.get()}。")
            else:
                self.status_var.set(f"解析完成，当前标准：{self.protocol_var.get()}。")
        except ParseError as exc:
            self.show_output(
                f"解析失败：{exc}\n\n"
                f"当前选择标准：{self.protocol_var.get()}\n"
                "建议：如果该帧在另一个标准下能够解析，请切换顶部“协议标准”后重新解析。"
            )
            self.status_var.set("解析失败。")
        except Exception as exc:
            self.show_output(f"未知错误：{exc}")
            self.status_var.set("解析失败。")

    def copy_output(self):
        content = self.output_text.get("1.0", tk.END).strip()
        if not content:
            messagebox.showinfo("提示", "当前没有可复制的解析结果。")
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(content)
        self.status_var.set("解析结果已复制到剪贴板。")

    def batch_parse_file(self):
        input_path = filedialog.askopenfilename(
            title="选择包含 AT+ZIPSEND 的日志文件",
            filetypes=[
                ("日志/文本文件", "*.txt *.log *.csv *.dat"),
                ("所有文件", "*.*"),
            ],
        )
        if not input_path:
            return

        input_stem = os.path.splitext(os.path.basename(input_path))[0]
        default_name = f"{input_stem}_{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
        output_path = filedialog.asksaveasfilename(
            title="保存解析结果 Excel",
            initialdir=os.path.dirname(input_path),
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
        )
        if not output_path:
            return

        try:
            self.status_var.set("正在批量解析文件，请稍等...")
            self.root.update_idletasks()
            summary = batch_parse_zipsend_file(input_path, output_path, self.protocol_var.get())
            self.status_var.set(
                f"批量解析完成：{summary['parsed_count']} 条，错误 {summary['error_count']} 条，0200位置 {summary['location_count']} 条。"
            )
            self.show_output(
                "批量解析完成\n"
                f"当前标准：{self.protocol_var.get()}\n"
                f"输入文件：{summary['input_path']}\n"
                f"输出文件：{summary['output_path']}\n"
                f"解析条数：{summary['parsed_count']}\n"
                f"错误条数：{summary['error_count']}\n"
                f"0200位置条数：{summary['location_count']}"
            )
            messagebox.showinfo(
                "批量解析完成",
                f"已生成 Excel：\n{summary['output_path']}\n\n"
                f"解析条数：{summary['parsed_count']}\n"
                f"错误条数：{summary['error_count']}\n"
                f"0200位置条数：{summary['location_count']}",
            )
        except Exception as exc:
            self.show_output(f"解析失败：{exc}")
            self.status_var.set("批量解析失败。")
            messagebox.showerror("批量解析失败", str(exc))


def main():
    root = tk.Tk()
    app = JT808ParserApp(root)
    root.mainloop()
    return app


if __name__ == "__main__":
    main()
