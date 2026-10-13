from __future__ import annotations

import os
import queue
import re
import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - only used on machines missing the dependency
    raise SystemExit(
        "缺少 openpyxl，请先在命令行执行：python -m pip install openpyxl"
    ) from exc


APP_TITLE = "ZB_TOPIC关键字时间差统计工具"
TIMESTAMP_RE = re.compile(
    r"^\s*(?P<timestamp>\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}(?:[.,]\d{1,6})?)"
)
ILLEGAL_EXCEL_CHARS_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
AUTO_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "gbk")
PREVIEW_LIMIT = 500
KEYWORD_SPLIT_RE = re.compile(r"[\r\n,，;；]+")


@dataclass(frozen=True)
class MatchRecord:
    line_number: int
    full_line: str
    timestamp: datetime
    delta_seconds: float | None


@dataclass(frozen=True)
class KeywordResult:
    keyword: str
    records: list[MatchRecord]
    invalid_timestamp_lines: int
    truncated_excel_lines: int = 0


@dataclass(frozen=True)
class AnalysisResult:
    keyword_results: list[KeywordResult]
    encoding: str

    @property
    def total_records(self) -> int:
        return sum(len(item.records) for item in self.keyword_results)

    @property
    def total_intervals(self) -> int:
        return sum(max(len(item.records) - 1, 0) for item in self.keyword_results)

    @property
    def total_invalid_timestamp_lines(self) -> int:
        return sum(item.invalid_timestamp_lines for item in self.keyword_results)

    @property
    def total_truncated_excel_lines(self) -> int:
        return sum(item.truncated_excel_lines for item in self.keyword_results)


def parse_keywords(keyword_text: str) -> list[str]:
    """按换行、中文/英文逗号或分号拆分，并保持输入顺序去重。"""
    keywords: list[str] = []
    seen: set[str] = set()
    for part in KEYWORD_SPLIT_RE.split(keyword_text):
        keyword = part.strip()
        if keyword and keyword not in seen:
            keywords.append(keyword)
            seen.add(keyword)
    return keywords


def parse_timestamp(line: str) -> datetime | None:
    """提取行首时间，支持秒以及最多 6 位小数秒。"""
    match = TIMESTAMP_RE.match(line)
    if not match:
        return None

    timestamp_text = match.group("timestamp").replace("T", " ").replace(",", ".")
    try:
        return datetime.fromisoformat(timestamp_text)
    except ValueError:
        return None


def _analyze_with_encoding(
    txt_path: Path, keywords: list[str], encoding: str
) -> list[KeywordResult]:
    records_by_keyword: dict[str, list[MatchRecord]] = {
        keyword: [] for keyword in keywords
    }
    invalid_counts = {keyword: 0 for keyword in keywords}
    previous_timestamps: dict[str, datetime | None] = {
        keyword: None for keyword in keywords
    }

    with txt_path.open("r", encoding=encoding, errors="strict", newline=None) as file:
        for line_number, raw_line in enumerate(file, start=1):
            matched_keywords = [keyword for keyword in keywords if keyword in raw_line]
            if not matched_keywords:
                continue

            full_line = raw_line.rstrip("\r\n")
            timestamp = parse_timestamp(full_line)
            if timestamp is None:
                for keyword in matched_keywords:
                    invalid_counts[keyword] += 1
                continue

            for keyword in matched_keywords:
                previous_timestamp = previous_timestamps[keyword]
                delta_seconds = None
                if previous_timestamp is not None:
                    delta_seconds = (timestamp - previous_timestamp).total_seconds()

                records_by_keyword[keyword].append(
                    MatchRecord(
                        line_number=line_number,
                        full_line=full_line,
                        timestamp=timestamp,
                        delta_seconds=delta_seconds,
                    )
                )
                previous_timestamps[keyword] = timestamp

    return [
        KeywordResult(
            keyword=keyword,
            records=records_by_keyword[keyword],
            invalid_timestamp_lines=invalid_counts[keyword],
        )
        for keyword in keywords
    ]


def analyze_txt_file(
    txt_path: str | Path,
    keywords: str | list[str] | tuple[str, ...],
    encoding_choice: str = "自动识别",
) -> AnalysisResult:
    path = Path(txt_path)
    if not path.is_file():
        raise FileNotFoundError(f"TXT文件不存在：{path}")

    if isinstance(keywords, str):
        keyword_list = parse_keywords(keywords)
    else:
        keyword_list = parse_keywords("\n".join(str(item) for item in keywords))
    if not keyword_list:
        raise ValueError("请至少输入一个关键字。")

    if encoding_choice == "自动识别":
        encodings = AUTO_ENCODINGS
    else:
        encoding_map = {
            "UTF-8": "utf-8-sig",
            "GB18030": "gb18030",
            "GBK": "gbk",
        }
        encodings = (encoding_map.get(encoding_choice, encoding_choice),)

    decode_errors: list[str] = []
    for encoding in encodings:
        try:
            keyword_results = _analyze_with_encoding(path, keyword_list, encoding)
            return AnalysisResult(
                keyword_results=keyword_results,
                encoding=encoding,
            )
        except UnicodeDecodeError as exc:
            decode_errors.append(f"{encoding}: {exc}")

    raise UnicodeError(
        "无法识别TXT文件编码。请在界面中手动选择编码。\n" + "\n".join(decode_errors)
    )


def sanitize_excel_text(text: str) -> tuple[str, bool]:
    text = ILLEGAL_EXCEL_CHARS_RE.sub("", text)
    if len(text) <= 32767:
        return text, False
    return text[:32764] + "...", True


def export_to_excel(
    result: AnalysisResult,
    output_path: str | Path,
    source_path: str | Path,
) -> AnalysisResult:
    destination = Path(output_path)
    if destination.suffix.lower() != ".xlsx":
        destination = destination.with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)

    headers = ("包含关键字的完整行", "与上一条记录时间差（秒）")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10, color="222222")
    thin_gray = Side(style="thin", color="D9E2F3")

    workbook = Workbook()
    exported_keyword_results: list[KeywordResult] = []
    for index, keyword_result in enumerate(result.keyword_results, start=1):
        sheet = workbook.active if index == 1 else workbook.create_sheet()
        sheet.title = f"Sheet{index}"
        # 保留 Excel 默认网格线，避免数据区域显示成白茫茫的一片。
        sheet.sheet_view.showGridLines = True
        sheet.freeze_panes = "A2"
        sheet.append(headers)

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin_gray)
        sheet.row_dimensions[1].height = 26

        truncated_count = 0
        for record in keyword_result.records:
            safe_line, was_truncated = sanitize_excel_text(record.full_line)
            truncated_count += int(was_truncated)
            sheet.append((safe_line, record.delta_seconds))

        for row in sheet.iter_rows(min_row=2, max_col=2):
            row[0].font = body_font
            row[0].alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=False
            )
            row[1].font = body_font
            row[1].alignment = Alignment(horizontal="right", vertical="center")
            row[1].number_format = "0.######"

        sheet.auto_filter.ref = f"A1:B{max(sheet.max_row, 1)}"
        sheet.column_dimensions[get_column_letter(1)].width = 105
        sheet.column_dimensions[get_column_letter(2)].width = 25
        exported_keyword_results.append(
            KeywordResult(
                keyword=keyword_result.keyword,
                records=keyword_result.records,
                invalid_timestamp_lines=keyword_result.invalid_timestamp_lines,
                truncated_excel_lines=truncated_count,
            )
        )

    workbook.properties.title = "ZB_TOPIC关键字时间差统计"
    workbook.properties.subject = f"关键字数量：{len(result.keyword_results)}"
    sheet_mapping_description = (
        f"来源文件：{Path(source_path).name}；"
        + "；".join(
            f"Sheet{index}={item.keyword}"
            for index, item in enumerate(result.keyword_results, start=1)
        )
    )
    workbook.properties.description = sheet_mapping_description[:2000]

    temp_path = destination.with_name(
        f".{destination.stem}.{os.getpid()}.{threading.get_ident()}.tmp.xlsx"
    )
    try:
        workbook.save(temp_path)
        os.replace(temp_path, destination)
    finally:
        workbook.close()
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass

    return AnalysisResult(
        keyword_results=exported_keyword_results,
        encoding=result.encoding,
    )


class TimeDifferenceApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("940x710")
        self.root.minsize(820, 600)

        self.source_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.encoding_var = tk.StringVar(value="自动识别")
        self.status_var = tk.StringVar(value="请选择TXT文件并输入一个或多个关键字。")
        self.summary_var = tk.StringVar(value="关键字：0    匹配记录：0    时间间隔：0")
        self.preview_sheet_var = tk.StringVar()
        self._message_queue: queue.Queue[tuple[str, object]] = queue.Queue()
        self._preview_results: dict[str, KeywordResult] = {}

        self._configure_style()
        self._build_ui()
        self.root.after(100, self._poll_worker_message)

    def _configure_style(self) -> None:
        style = ttk.Style(self.root)
        try:
            style.theme_use("vista")
        except tk.TclError:
            pass
        style.configure("TLabel", font=("Microsoft YaHei UI", 10))
        style.configure("TButton", font=("Microsoft YaHei UI", 10), padding=(10, 5))
        style.configure("TEntry", font=("Microsoft YaHei UI", 10))
        style.configure("Treeview", font=("Microsoft YaHei UI", 9), rowheight=27)
        style.configure("Treeview.Heading", font=("Microsoft YaHei UI", 10, "bold"))
        style.configure("Accent.TButton", font=("Microsoft YaHei UI", 10, "bold"))

    def _build_ui(self) -> None:
        container = ttk.Frame(self.root, padding=18)
        container.pack(fill="both", expand=True)
        container.columnconfigure(1, weight=1)
        container.rowconfigure(5, weight=1)

        title = ttk.Label(
            container,
            text="ZB_TOPIC关键字时间差统计",
            font=("Microsoft YaHei UI", 17, "bold"),
            foreground="#1F4E78",
        )
        title.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 14))

        ttk.Label(container, text="TXT文件：").grid(row=1, column=0, sticky="e", pady=5)
        self.source_entry = ttk.Entry(container, textvariable=self.source_var)
        self.source_entry.grid(row=1, column=1, sticky="ew", padx=(0, 8), pady=5)
        self.select_source_button = ttk.Button(
            container, text="选择文件", command=self._select_source
        )
        self.select_source_button.grid(row=1, column=2, sticky="ew", pady=5)

        ttk.Label(container, text="关键字：").grid(row=2, column=0, sticky="ne", pady=5)
        keyword_frame = ttk.Frame(container)
        keyword_frame.grid(row=2, column=1, sticky="ew", padx=(0, 8), pady=5)
        keyword_frame.columnconfigure(0, weight=1)
        self.keyword_text = tk.Text(
            keyword_frame,
            height=4,
            wrap="none",
            font=("Microsoft YaHei UI", 10),
            relief="solid",
            borderwidth=1,
            highlightthickness=0,
            padx=6,
            pady=5,
        )
        self.keyword_text.grid(row=0, column=0, sticky="ew")
        self.keyword_text.insert("1.0", "loc/862710080532022")
        ttk.Label(
            keyword_frame,
            text="建议每行输入一个关键字；也支持英文/中文逗号或分号分隔。重复关键字会自动去重。",
            foreground="#666666",
        ).grid(row=1, column=0, sticky="w", pady=(4, 0))
        self.encoding_combo = ttk.Combobox(
            container,
            textvariable=self.encoding_var,
            values=("自动识别", "UTF-8", "GB18030", "GBK"),
            state="readonly",
            width=12,
        )
        self.encoding_combo.grid(row=2, column=2, sticky="ew", pady=5)

        ttk.Label(container, text="Excel文件：").grid(row=3, column=0, sticky="e", pady=5)
        self.output_entry = ttk.Entry(container, textvariable=self.output_var)
        self.output_entry.grid(row=3, column=1, sticky="ew", padx=(0, 8), pady=5)
        self.select_output_button = ttk.Button(
            container, text="保存位置", command=self._select_output
        )
        self.select_output_button.grid(row=3, column=2, sticky="ew", pady=5)

        action_frame = ttk.Frame(container)
        action_frame.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(12, 10))
        action_frame.columnconfigure(0, weight=1)
        ttk.Label(action_frame, textvariable=self.summary_var).grid(row=0, column=0, sticky="w")
        self.run_button = ttk.Button(
            action_frame,
            text="分析并导出Excel",
            style="Accent.TButton",
            command=self._start_analysis,
        )
        self.run_button.grid(row=0, column=1, sticky="e")

        preview_frame = ttk.LabelFrame(container, text="结果预览", padding=8)
        preview_frame.grid(row=5, column=0, columnspan=3, sticky="nsew")
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(1, weight=1)

        preview_toolbar = ttk.Frame(preview_frame)
        preview_toolbar.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 7))
        preview_toolbar.columnconfigure(1, weight=1)
        ttk.Label(preview_toolbar, text="预览工作表：").grid(row=0, column=0, sticky="w")
        self.preview_sheet_combo = ttk.Combobox(
            preview_toolbar,
            textvariable=self.preview_sheet_var,
            state="disabled",
        )
        self.preview_sheet_combo.grid(row=0, column=1, sticky="ew")
        self.preview_sheet_combo.bind(
            "<<ComboboxSelected>>", self._on_preview_sheet_changed
        )

        self.tree = ttk.Treeview(
            preview_frame,
            columns=("line", "delta"),
            show="headings",
            selectmode="browse",
        )
        self.tree.heading("line", text="包含关键字的完整行")
        self.tree.heading("delta", text="与上一条记录时间差（秒）")
        self.tree.column("line", width=680, minwidth=350, anchor="w")
        self.tree.column("delta", width=185, minwidth=160, anchor="e", stretch=False)

        vertical_scrollbar = ttk.Scrollbar(
            preview_frame, orient="vertical", command=self.tree.yview
        )
        horizontal_scrollbar = ttk.Scrollbar(
            preview_frame, orient="horizontal", command=self.tree.xview
        )
        self.tree.configure(
            yscrollcommand=vertical_scrollbar.set,
            xscrollcommand=horizontal_scrollbar.set,
        )
        self.tree.grid(row=1, column=0, sticky="nsew")
        vertical_scrollbar.grid(row=1, column=1, sticky="ns")
        horizontal_scrollbar.grid(row=2, column=0, sticky="ew")

        status_frame = ttk.Frame(container)
        status_frame.grid(row=6, column=0, columnspan=3, sticky="ew", pady=(10, 0))
        status_frame.columnconfigure(0, weight=1)
        self.status_label = ttk.Label(
            status_frame,
            textvariable=self.status_var,
            foreground="#555555",
            wraplength=700,
        )
        self.status_label.grid(row=0, column=0, sticky="w")
        self.progress = ttk.Progressbar(status_frame, mode="indeterminate", length=150)
        self.progress.grid(row=0, column=1, sticky="e")
        self.progress.grid_remove()

    def _select_source(self) -> None:
        selected = filedialog.askopenfilename(
            title="选择TXT文件",
            filetypes=(("文本文件", "*.txt"), ("日志文件", "*.log"), ("所有文件", "*.*")),
        )
        if not selected:
            return
        self.source_var.set(selected)
        source_path = Path(selected)
        self.output_var.set(str(source_path.with_name(f"{source_path.stem}_多关键字时间差.xlsx")))

    def _select_output(self) -> None:
        initial_path = Path(self.output_var.get()) if self.output_var.get().strip() else None
        selected = filedialog.asksaveasfilename(
            title="选择Excel保存位置",
            defaultextension=".xlsx",
            filetypes=(("Excel工作簿", "*.xlsx"),),
            initialdir=str(initial_path.parent) if initial_path else None,
            initialfile=initial_path.name if initial_path else "多关键字时间差.xlsx",
        )
        if selected:
            self.output_var.set(selected)

    def _start_analysis(self) -> None:
        source = self.source_var.get().strip().strip('"')
        keyword_text = self.keyword_text.get("1.0", "end").strip()
        keywords = parse_keywords(keyword_text)
        output = self.output_var.get().strip().strip('"')

        if not source:
            messagebox.showwarning("提示", "请先选择TXT文件。")
            return
        if not Path(source).is_file():
            messagebox.showerror("错误", "选择的TXT文件不存在。")
            return
        if not keywords:
            messagebox.showwarning("提示", "请至少输入一个要统计的关键字。")
            return
        if not output:
            source_path = Path(source)
            output = str(source_path.with_name(f"{source_path.stem}_多关键字时间差.xlsx"))
            self.output_var.set(output)

        if Path(output).suffix.lower() != ".xlsx":
            output = str(Path(output).with_suffix(".xlsx"))
            self.output_var.set(output)

        self._set_busy(True)
        self.status_var.set(f"正在统计 {len(keywords)} 个关键字并生成Excel……")
        worker = threading.Thread(
            target=self._worker,
            args=(source, keywords, output, self.encoding_var.get()),
            daemon=True,
        )
        worker.start()

    def _worker(
        self, source: str, keywords: list[str], output: str, encoding: str
    ) -> None:
        try:
            result = analyze_txt_file(source, keywords, encoding)
            exported_result = export_to_excel(result, output, source)
            self._message_queue.put(("success", (exported_result, output)))
        except Exception as exc:  # GUI displays the concrete error instead of terminating
            self._message_queue.put(("error", exc))

    def _poll_worker_message(self) -> None:
        try:
            message_type, payload = self._message_queue.get_nowait()
        except queue.Empty:
            self.root.after(100, self._poll_worker_message)
            return

        self._set_busy(False)
        if message_type == "success":
            result, output = payload
            self._show_result(result)
            notes = []
            if result.total_invalid_timestamp_lines:
                notes.append(
                    f"跳过 {result.total_invalid_timestamp_lines} 条匹配（关键字存在，但行首时间无效）"
                )
            if result.total_truncated_excel_lines:
                notes.append(
                    f"{result.total_truncated_excel_lines} 行超过Excel单元格长度限制，已截断"
                )
            note_text = "；".join(notes)
            self.status_var.set(
                f"完成。编码：{result.encoding}。Excel已保存：{output}"
                + (f"。{note_text}" if note_text else "")
            )
            mapping_lines = [
                f"Sheet{index}：{item.keyword}（{len(item.records)} 条）"
                for index, item in enumerate(result.keyword_results, start=1)
            ]
            if len(mapping_lines) > 12:
                mapping_lines = mapping_lines[:12] + [
                    f"……其余 {len(result.keyword_results) - 12} 个工作表请在Excel中查看"
                ]
            sheet_mapping = "\n".join(mapping_lines)
            messagebox.showinfo(
                "导出完成",
                f"共统计 {len(result.keyword_results)} 个关键字，找到 {result.total_records} 条记录，"
                f"计算出 {result.total_intervals} 个时间间隔。\n\n"
                f"{sheet_mapping}\n\nExcel文件：\n{output}",
            )
        else:
            error = payload
            self.status_var.set(f"处理失败：{error}")
            messagebox.showerror("处理失败", str(error))

        self.root.after(100, self._poll_worker_message)

    def _show_result(self, result: AnalysisResult) -> None:
        self._preview_results.clear()
        selector_values = []
        for index, keyword_result in enumerate(result.keyword_results, start=1):
            label = f"Sheet{index} - {keyword_result.keyword}"
            selector_values.append(label)
            self._preview_results[label] = keyword_result

        self.preview_sheet_combo.configure(values=selector_values, state="readonly")
        if selector_values:
            self.preview_sheet_var.set(selector_values[0])
            self._show_keyword_result(self._preview_results[selector_values[0]])
        else:
            self.preview_sheet_var.set("")
            self._clear_tree()

        self.summary_var.set(
            f"关键字：{len(result.keyword_results)}    匹配记录：{result.total_records}    "
            f"时间间隔：{result.total_intervals}"
        )

    def _show_keyword_result(self, result: KeywordResult) -> None:
        self._clear_tree()
        for record in result.records[:PREVIEW_LIMIT]:
            delta_text = "" if record.delta_seconds is None else f"{record.delta_seconds:g}"
            self.tree.insert("", "end", values=(record.full_line, delta_text))

    def _on_preview_sheet_changed(self, _event: object = None) -> None:
        selected = self.preview_sheet_var.get()
        keyword_result = self._preview_results.get(selected)
        if keyword_result is not None:
            self._show_keyword_result(keyword_result)

    def _clear_tree(self) -> None:
        for item_id in self.tree.get_children():
            self.tree.delete(item_id)

    def _set_busy(self, busy: bool) -> None:
        state = "disabled" if busy else "normal"
        self.run_button.configure(state=state)
        self.select_source_button.configure(state=state)
        self.select_output_button.configure(state=state)
        self.source_entry.configure(state=state)
        self.keyword_text.configure(state=state)
        self.encoding_combo.configure(state="disabled" if busy else "readonly")
        self.output_entry.configure(state=state)
        self.preview_sheet_combo.configure(
            state="disabled" if busy or not self._preview_results else "readonly"
        )
        if busy:
            self.progress.configure(value=0)
            self.progress.grid()
            self.progress.start(10)
        else:
            self.progress.stop()
            self.progress.configure(value=0)
            self.progress.grid_remove()


def main() -> None:
    root = tk.Tk()
    TimeDifferenceApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
