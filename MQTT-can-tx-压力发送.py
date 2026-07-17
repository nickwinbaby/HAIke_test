import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
import threading
import time
import re
from openpyxl import load_workbook
import paho.mqtt.client as mqtt


class MQTTHexPressureTool:

    def __init__(self, root):
        self.root = root
        self.root.title("MQTT HEX压测工具 V1.0 - Mangetic - 260609 ★★★★★ 广告位招租 ★★★★★")
        self.root.geometry("900x700")

        self.center_window(900, 700)

        self.client = None
        self.excel_path = ""

        self.stop_flag = False

        self.success = 0
        self.fail = 0

        self.build_ui()

    # =========================
    # 窗口居中
    # =========================
    def center_window(self, w, h):
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        x = int((sw - w) / 2)
        y = int((sh - h) / 2)
        self.root.geometry(f"{w}x{h}+{x}+{y}")

    # =========================
    # UI
    # =========================
    def build_ui(self):

        label_w = 18
        entry_w = 55

        row = 0

        tk.Label(self.root, text="MQTT服务器：", width=label_w, anchor="e").grid(row=row, column=0)
        self.host = tk.Entry(self.root, width=entry_w)
        self.host.insert(0, "richtest.richlinkcloud.com")
        self.host.grid(row=row, column=1, padx=5, pady=3)

        row += 1

        tk.Label(self.root, text="端口：", width=label_w, anchor="e").grid(row=row, column=0)
        self.port = tk.Entry(self.root, width=entry_w)
        self.port.insert(0, "1883")
        self.port.grid(row=row, column=1)

        row += 1

        tk.Label(self.root, text="用户名：", width=label_w, anchor="e").grid(row=row, column=0)
        self.user = tk.Entry(self.root, width=entry_w)
        self.user.grid(row=row, column=1)

        row += 1

        tk.Label(self.root, text="密码：", width=label_w, anchor="e").grid(row=row, column=0)
        self.pwd = tk.Entry(self.root, width=entry_w, show="*")
        self.pwd.grid(row=row, column=1)

        row += 1

        tk.Label(self.root, text="Topic：", width=label_w, anchor="e").grid(row=row, column=0)
        self.topic = tk.Entry(self.root, width=entry_w)
        self.topic.grid(row=row, column=1)

        row += 1

        tk.Label(self.root, text="单条间隔(秒)：", width=label_w, anchor="e").grid(row=row, column=0)
        self.interval = tk.Entry(self.root, width=entry_w)
        self.interval.insert(0, "0")
        self.interval.grid(row=row, column=1)

        row += 1

        tk.Label(self.root, text="循环间隔(秒)：", width=label_w, anchor="e").grid(row=row, column=0)
        self.loop_interval = tk.Entry(self.root, width=entry_w)
        self.loop_interval.insert(0, "30")
        self.loop_interval.grid(row=row, column=1)

        row += 1

        tk.Label(self.root, text="重复次数：", width=label_w, anchor="e").grid(row=row, column=0)
        self.repeat = tk.Entry(self.root, width=entry_w)
        self.repeat.insert(0, "5")
        self.repeat.grid(row=row, column=1)

        row += 1

        tk.Button(self.root, text="连接MQTT", command=self.connect_mqtt, width=15).grid(row=row, column=0, pady=5)
        tk.Button(self.root, text="选择Excel", command=self.select_excel, width=15).grid(row=row, column=1)

        row += 1

        self.file_label = tk.Label(self.root, text="未选择Excel文件")
        self.file_label.grid(row=row, column=0, columnspan=2)

        row += 1

        tk.Button(self.root, text="开始发送", bg="green", fg="white",
                  command=self.start, width=20).grid(row=row, column=0, pady=10)

        tk.Button(self.root, text="停止发送", bg="red", fg="white",
                  command=self.stop, width=20).grid(row=row, column=1)

        row += 1

        self.stat = tk.Label(self.root, text="成功: 0   失败: 0")
        self.stat.grid(row=row, column=0, columnspan=2)

        row += 1

        self.log_box = scrolledtext.ScrolledText(self.root, width=110, height=25)
        self.log_box.grid(row=row, column=0, columnspan=2, padx=5, pady=10)

    # =========================
    # log
    # =========================
    def log(self, msg):
        t = time.strftime("%H:%M:%S")
        self.log_box.insert(tk.END, f"[{t}] {msg}\n")
        self.log_box.see(tk.END)

    # =========================
    # MQTT连接（兼容1.6.1）
    # =========================
    def connect_mqtt(self):

        try:
            host = self.host.get().strip()
            port = int(self.port.get().strip())
            user = self.user.get().strip()
            pwd = self.pwd.get().strip()

            self.client = mqtt.Client()

            if user:
                self.client.username_pw_set(user, pwd)

            self.client.on_connect = self.on_connect

            self.client.connect(host, port, 60)
            self.client.loop_start()

            self.log(f"连接MQTT中：{host}:{port}")

        except Exception as e:
            messagebox.showerror("错误", str(e))

    def on_connect(self, client, userdata, flags, rc):
        if rc == 0:
            self.log("MQTT连接成功")
        else:
            self.log(f"MQTT连接失败：{rc}")

    # =========================
    # Excel
    # =========================
    def select_excel(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if path:
            self.excel_path = path
            self.file_label.config(text=path)

    # =========================
    # HEX校验
    # =========================
    def is_hex(self, s):
        return bool(re.fullmatch(r"[0-9A-Fa-f]+", s))

    # =========================
    # start
    # =========================
    def start(self):

        if not self.client:
            messagebox.showwarning("提示", "请先连接MQTT")
            return

        if not self.excel_path:
            messagebox.showwarning("提示", "请选择Excel")
            return

        self.stop_flag = False

        threading.Thread(target=self.run, daemon=True).start()

    def stop(self):
        self.stop_flag = True
        self.log("停止发送")

    # =========================
    # 核心发送逻辑
    # =========================
    def run(self):

        topic = self.topic.get().strip()
        interval = float(self.interval.get())
        loop_interval = float(self.loop_interval.get())
        repeat = int(self.repeat.get())

        wb = load_workbook(self.excel_path)
        ws = wb.active

        for r in range(repeat):

            if self.stop_flag:
                break

            self.log(f"====== 第{r+1}轮开始 ======")

            row = 1

            while True:

                if self.stop_flag:
                    break

                v = ws.cell(row=row, column=1).value

                if v is None:
                    break

                data = str(v).strip()

                if not data:
                    break

                if not self.is_hex(data):
                    self.fail += 1
                    self.log(f"非法HEX(Row {row}): {data}")
                    row += 1
                    continue

                try:
                    payload = bytes.fromhex(data)

                    res = self.client.publish(topic, payload, qos=1)

                    if res.rc == 0:
                        self.success += 1
                        self.log(f"发送成功 Row={row} Len={len(payload)}")
                    else:
                        self.fail += 1
                        self.log(f"发送失败 Row={row}")

                except Exception as e:
                    self.fail += 1
                    self.log(f"异常 Row={row}: {e}")

                self.stat.config(text=f"成功: {self.success}   失败: {self.fail}")

                row += 1

                if interval > 0:
                    time.sleep(interval)

            self.log(f"====== 第{r+1}轮结束 ======")

            if r != repeat - 1:
                self.log(f"等待 {loop_interval}s ...")
                time.sleep(loop_interval)

        self.log("全部完成")


if __name__ == "__main__":
    root = tk.Tk()
    app = MQTTHexPressureTool(root)
    root.mainloop()