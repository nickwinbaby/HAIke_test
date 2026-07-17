import tkinter as tk
from tkinter import messagebox, scrolledtext
import requests
import time
import os
import json
from datetime import datetime, timedelta, timezone
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font, GradientFill
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.font_manager import FontProperties

def main(imei_num_list: list[str], start_time: str, stop_time: str):
    file_path = 'C:/SENSOR-专项测试/数据检查/'
    os.makedirs(file_path, exist_ok=True)
    current_time = datetime.now().strftime('%Y%m%d_%H%M%S')

    start_dt = datetime.strptime(start_time, "%Y-%m-%d %H:%M:%S")
    stop_dt = datetime.strptime(stop_time, "%Y-%m-%d %H:%M:%S")

    start_timestamp = int(start_dt.replace(tzinfo=timezone.utc).timestamp() * 1000)
    stop_timestamp = int(stop_dt.replace(tzinfo=timezone.utc).timestamp() * 1000)

    def convert_to_beijing_time(timestamp):
        dt_utc = datetime.fromtimestamp(timestamp / 1000, tz=timezone.utc)
        dt_beijing = dt_utc + timedelta(hours=8)
        return dt_beijing.strftime('%Y-%m-%d %H:%M:%S')

    start_time_beijing = convert_to_beijing_time(start_timestamp)
    stop_time_beijing = convert_to_beijing_time(stop_timestamp)
    data_request_time = f"数据开始时间:{start_time_beijing}", f"数据结束时间: {stop_time_beijing}"

    wb = Workbook()
    ws = wb.active

    blue_fill = PatternFill(start_color="6d9eeb", end_color="6d9eeb", fill_type="solid")
    light_yellow_fill = PatternFill(start_color="f9ca9c", end_color="f9ca9c", fill_type="solid")

    headers = ['基本信息', '最后一包数据', '急加速', '急减速', '急转弯', '疑似碰撞', '休眠时间', '唤醒时间',
               '点火时间', '熄火时间', 'SENSOR自学习状态', '行程统计', '自学习次数', '学习成功率', '行驶里程',
               '重启事件', '备注']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = blue_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # 省略 parse_event、hex_to_beijing_time、server_get_data 的完整逻辑
    # 直接使用你原来的逻辑
    # 这里 server_get_data 内部会生成 Excel 和 JSON

    # ↓ 调用内部函数获取数据
    def server_get_data(imei_num_list, start_timestamp, stop_timestamp):
        # 这里复制你原来的 server_get_data 内部逻辑
        pass  # 实际使用时，把你原来的函数逻辑直接放进来

    server_get_data(imei_num_list, start_timestamp, stop_timestamp)

    ws['Q2'] = f"{data_request_time[0]}\n{data_request_time[1]}"
    ws['Q2'].fill = light_yellow_fill

    # 列宽设置、字体颜色、保存 Excel 等逻辑照原来写
    # filename = ...
    # wb.save(filename)
    # os.startfile(file_path)

def create_gui():
    window = tk.Tk()
    window.title("SENSOR专项-服务器获取数据-木林夕")
    window.geometry("500x400")

    tk.Label(window, text="IMEI 列表（每行一个）:").pack(pady=(10, 0))
    imei_text = scrolledtext.ScrolledText(window, width=60, height=10)
    imei_text.pack(pady=5)

    tk.Label(window, text="开始时间-输入无时区时间 (YYYY-MM-DD HH:MM:SS):").pack()
    start_entry = tk.Entry(window, width=30)
    start_entry.insert(0, "2025-05-19 01:00:00")
    start_entry.pack(pady=2)

    tk.Label(window, text="结束时间-输入无时区时间 (YYYY-MM-DD HH:MM:SS):").pack()
    end_entry = tk.Entry(window, width=30)
    end_entry.insert(0, "2025-05-20 01:00:00")
    end_entry.pack(pady=2)

    def on_submit():
        imei_list_raw = imei_text.get("1.0", tk.END).strip()
        imeis = [line.strip() for line in imei_list_raw.splitlines() if line.strip()]
        start_time = start_entry.get().strip()
        stop_time = end_entry.get().strip()

        if not imeis or not start_time or not stop_time:
            messagebox.showerror("错误", "请填写所有字段！")
            return

        try:
            main(imeis, start_time, stop_time)
            messagebox.showinfo("完成", f"已成功处理 {len(imeis)} 个 IMEI。")
        except Exception as e:
            messagebox.showerror("执行出错", str(e))

    tk.Button(window, text="运行程序", command=on_submit, bg="green", fg="white").pack(pady=20)
    window.mainloop()
