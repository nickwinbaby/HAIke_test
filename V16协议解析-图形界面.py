import tkinter as tk
from tkinter import filedialog, messagebox
import re
import os
from openpyxl import Workbook
from datetime import datetime
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def remove_bracket_content(text):
    text = re.sub(r'\[.*?\]', '', text)
    return text.replace(' ', '')

def auto_adjust_column_width(ws, row_num):
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_num, column=col)
        max_length = len(str(cell.value)) if cell.value else 0
        adjusted_width = (max_length + 2)
        ws.column_dimensions[get_column_letter(col)].width = adjusted_width

def process_file(input_file, output_dir):
    current_time = datetime.now().strftime('%Y%m%d_%H%M%S')
    excel_filename = f'{output_dir}/{current_time}_V16.xlsx'

    wb = Workbook()
    ws = wb.active

    light_grey_fill = PatternFill(start_color="c8c7c7", end_color="c8c7c7", fill_type="solid")
    light_blue_fill = PatternFill(start_color="6d9eeb", end_color="6d9eeb", fill_type="solid")
    light_yellow_fill = PatternFill(start_color="f9ca9c", end_color="f9ca9c", fill_type="solid")
    light_green_fill = PatternFill(start_color="b6d7a8", end_color="b6d7a8", fill_type="solid")
    light_red_fill = PatternFill(start_color="e99998", end_color="e99998", fill_type="solid")

    ws.merge_cells('B1:E1')
    ws.merge_cells('F1:I1')
    ws.merge_cells('J1:R1')
    ws.merge_cells('S1:AA1')

    ws['B1'] = 'MESSAGE'
    ws['F1'] = 'DEVICE'
    ws['J1'] = 'QUALITY'
    ws['S1'] = 'LOCATION'

    ws['A2'] = 'RAW DATA'
    ws['B2'] = 'PACKET LENGTH'
    ws['C2'] = 'PACKET VERSION'
    ws['D2'] = "PACKET TYPE"
    ws['E2'] = "SEQUENCE"
    ws['F2'] = "MANUFACTURER ID"
    ws['G2'] = "SOFTWARE VERSION"
    ws['H2'] = "HARDWARE VERSION"
    ws['I2'] = "DEVICE ID"
    ws['J2'] = "BATTERY %"
    ws['K2'] = "ACTIVE TIME"
    ws['L2'] = "IMEI"
    ws['M2'] = "CELL INFO (ECL)"
    ws['N2'] = "RSSI"
    ws['O2'] = "RSRP"
    ws['P2'] = "RSRQ"
    ws['Q2'] = "PLMN"
    ws['R2'] = "AUX"
    ws['S2'] = "GPS N/S"
    ws['T2'] = "GPS LATITUDE"
    ws['U2'] = "GPS E/W"
    ws['V2'] = "GPS LONGITUDE"
    ws['W2'] = "GPS TIMESTAMP (YYYYMMDDHHMISS)"
    ws['X2'] = "GPS ALTITUDE"
    ws['Y2'] = "GPS HORIZ EPEL"
    ws['Z2'] = "GPS SATELLITE COUNT"
    ws['AA2'] = "GPS HDOP"

    bold_font = Font(bold=True)
    center_alignment = Alignment(horizontal="center", vertical="center")

    def fill_cells(ws_state, cell_range, fill_color, font=None, alignment=None):
        for cell in cell_range:
            ws_state[cell].fill = fill_color
            if font:
                ws_state[cell].font = font
            if alignment:
                ws_state[cell].alignment = alignment

    fill_cells(ws, ['A1', 'A2'], light_grey_fill, bold_font, center_alignment)
    fill_cells(ws, ['B2', 'C2', 'D2', 'E2', 'B1'], light_blue_fill, bold_font, center_alignment)
    fill_cells(ws, ['F2', 'G2', 'H2', 'I2', 'F1'], light_yellow_fill, bold_font, center_alignment)
    fill_cells(ws, ['J2', 'K2', 'L2', 'M2', 'N2', 'O2', 'P2', 'Q2', 'R2', 'J1'], light_green_fill, bold_font, center_alignment)
    fill_cells(ws, ['S2', 'T2', 'U2', 'V2', 'W2', 'X2', 'Y2', 'Z2', 'AA2', 'S1'], light_red_fill, bold_font, center_alignment)

    auto_adjust_column_width(ws, 2)

    with open(input_file, 'r', encoding='utf-8') as file:
        row = 3
        for line in file:
            if "[TGAT]><-- UDP socket0 pack_len" in line:
                start_index = line.find('"') + 1
                end_index = line.find('"', start_index)
                if 0 < start_index < end_index:
                    content = line[start_index:end_index]
                    content = remove_bracket_content(content)
                    ws.cell(row=row, column=1, value=content).fill = light_grey_fill
                    ws.cell(row=row, column=2, value=content[0:3]).fill = light_blue_fill
                    ws.cell(row=row, column=3, value=content[3:6]).fill = light_blue_fill
                    ws.cell(row=row, column=4, value=content[6:7]).fill = light_blue_fill
                    ws.cell(row=row, column=5, value=content[7:10]).fill = light_blue_fill
                    ws.cell(row=row, column=6, value=content[10:14]).fill = light_yellow_fill
                    ws.cell(row=row, column=7, value=content[14:16]).fill = light_yellow_fill
                    ws.cell(row=row, column=8, value=content[16:18]).fill = light_yellow_fill
                    ws.cell(row=row, column=9, value=content[18:26]).fill = light_yellow_fill
                    ws.cell(row=row, column=10, value=content[26:28]).fill = light_green_fill
                    ws.cell(row=row, column=11, value=content[28:31]).fill = light_green_fill
                    ws.cell(row=row, column=12, value=content[31:46]).fill = light_green_fill
                    ws.cell(row=row, column=13, value=content[46:54]).fill = light_green_fill
                    ws.cell(row=row, column=14, value=content[54:58]).fill = light_green_fill
                    ws.cell(row=row, column=15, value=content[58:62]).fill = light_green_fill
                    ws.cell(row=row, column=16, value=content[62:66]).fill = light_green_fill
                    ws.cell(row=row, column=17, value=content[66:72]).fill = light_green_fill
                    ws.cell(row=row, column=18, value=content[72:77]).fill = light_green_fill
                    value = content[77:78]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=19, value=value).fill = light_red_fill
                    value = content[78:87]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=20, value=value).fill = light_red_fill
                    value = content[87:88]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=21, value=value).fill = light_red_fill
                    value = content[88:98]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=22, value=value).fill = light_red_fill
                    value = content[98:112]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=23, value=value).fill = light_red_fill
                    value = content[112:116]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=24, value=value).fill = light_red_fill
                    value = content[116:118]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=25, value=value).fill = light_red_fill
                    value = content[118:120]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=26, value=value).fill = light_red_fill
                    value = content[120:125]
                    if not value:
                        value = "NULL"
                    ws.cell(row=row, column=27, value=value).fill = light_red_fill
                    row += 1

    wb.save(excel_filename)
    os.startfile(output_dir)
    messagebox.showinfo("Success", "File processed successfully!")

def select_input_file():
    file_path = filedialog.askopenfilename(title="选择文件", filetypes=[("Text files", "*.TXT"), ("All files", "*.*")])
    input_file_var.set(file_path)

def select_output_dir():
    directory = filedialog.askdirectory(title="选择输出文件夹")
    output_dir_var.set(directory)

def start_processing():
    input_file = input_file_var.get()
    output_dir = output_dir_var.get()
    if not input_file or not output_dir:
        messagebox.showwarning("Warning", "Please select both input file and output directory.")
        return
    process_file(input_file, output_dir)

root = tk.Tk()
root.title("能能出品，必属精品-科旭定制版：2025-3-18")

input_file_var = tk.StringVar()
output_dir_var = tk.StringVar()

tk.Label(root, text="Input File:").grid(row=0, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=input_file_var, width=50).grid(row=0, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_input_file).grid(row=0, column=2, padx=10, pady=10)

tk.Label(root, text="Output Directory:").grid(row=1, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=output_dir_var, width=50).grid(row=1, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_output_dir).grid(row=1, column=2, padx=10, pady=10)

tk.Button(root, text="Start Processing", command=start_processing).grid(row=2, column=0, columnspan=3, pady=20)

root.mainloop()